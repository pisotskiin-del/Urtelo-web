"""Тесты карты абонента на выдуманных записях (без данных из DBF)."""

import io
import re
from collections import Counter
from datetime import date
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.management import call_command
from django.test import TestCase

from abonents import layout, reports
from catalogs.models import Akcept, Danrasch, GvsKot, Kategorpot, Limit, Minist, Potreb, SpTeplpoteri


class KartaBase(TestCase):
    def setUp(self):
        call_command("init_roles", stdout=open("/dev/null", "w"))
        self.ryb = Danrasch.objects.create(cod_gorod=1, name_gorod="Рыбница", name_mes="Октябрь", god=2026,
                                           tkoef=Decimal("0.229"), t_vozd=Decimal("10.0"))
        self.dub = Danrasch.objects.create(cod_gorod=2, name_gorod="Дубоссары", name_mes="Октябрь", god=2026)
        Minist.objects.create(cod_minist=1, name_minis="Без министерства")
        Kategorpot.objects.create(id_kotegor=1, name_kateg="Хозрасчет")
        GvsKot.objects.create(id_kot=16, name_kot="(Р) Котельная №2", kod_gor=self.ryb)
        Akcept.objects.create(id_akc=1, name_akc="БА")
        Akcept.objects.create(id_akc=3, name_akc="С")
        common = dict(cod_gor=self.ryb, cod_minist_id=1, id_kategor=1, id_koteln=16, accept=3, tip_rasot=1)
        self.a = Potreb.objects.create(potreb_id=5, nom_dog=6, name_potre='ООО "Крипс"', glav_potre="Магазин",
                                       id_katabon=1, nac_vsego=Decimal("10.5"), posle_prib=7, inf="Примечание",
                                       pok_new=Decimal("12.5"), dn_pereras=0, tkf_perer=Decimal("1.1"),
                                       gkl_perer=Decimal("3.3"), **common)
        self.b = Potreb.objects.create(potreb_id=7, nom_dog=6, name_potre='ООО "Крипс"', glav_potre="Склад",
                                       id_katabon=3, nom_dogab=6, nac_vsego=Decimal("2"), **common)
        self.c = Potreb.objects.create(potreb_id=9, nom_dog=8, name_potre="Школа  № 1", **common)
        Potreb.objects.create(potreb_id=11, nom_dog=8, name_potre="Чужой город", cod_gor=self.dub)

    def login(self, role):
        user = get_user_model().objects.create_user(role, password="x")
        user.groups.add(Group.objects.get(name=role))
        self.client.force_login(user)


class KartaTests(KartaBase):
    def test_each_editable_field_has_one_input_on_its_page(self):
        for key in layout.TAB_KEYS:
            _, _, groups = layout.tab(key)
            items = layout.HEADER + [f for _, fields in groups for f in fields]
            editable = Counter(f.name for f in items if not f.ro and not f.related)
            self.assertTrue(all(n == 1 for n in editable.values()), (key, editable))
            for name in editable:
                Potreb._meta.get_field(name)

    def test_viewer_sees_card_but_cannot_edit(self):
        self.login("Наблюдатель")
        response = self.client.get("/abonents/5/")
        self.assertContains(response, "Крипс")
        self.assertContains(response, "Рыбница")
        self.assertContains(response, "Октябрь 2026")
        self.assertContains(response, "(Р) Котельная №2")
        self.assertContains(response, 'href="/abonents/7/"')  # присоединённый прибор — ссылка
        self.assertNotContains(response, "Изменить")
        self.assertEqual(self.client.get("/abonents/5/edit/").status_code, 403)
        for key in layout.TAB_KEYS:
            self.assertEqual(self.client.get(f"/abonents/5/?tab={key}").status_code, 200)

    def test_navigation_stays_in_city(self):
        self.login("Наблюдатель")
        response = self.client.get("/abonents/9/")
        self.assertContains(response, "3 из 3")
        self.assertContains(response, 'href="/abonents/7/?tab=otop"')
        self.assertNotContains(response, "/abonents/11/")

    def test_saving_unchanged_page_keeps_all_values(self):
        self.login("Оператор")
        before = {f.attname: getattr(self.a, f.attname) for f in Potreb._meta.concrete_fields}
        for key in layout.TAB_KEYS:
            html = self.client.get(f"/abonents/5/edit/?tab={key}").content.decode()
            names = re.findall(r'name="([a-z_0-9]+)"', html)
            data = {}
            for name in set(names) - {"csrfmiddlewaretoken", "tab"}:
                self.assertEqual(names.count(name), 1, name)
                m = re.search(rf'<input[^>]*name="{name}"[^>]*value="([^"]*)"', html)
                if m:
                    data[name] = m.group(1).replace("&quot;", '"')
                    continue
                m = re.search(rf'<select name="{name}".*?<option value="([^"]*)"[^>]*selected', html, re.S)
                if m:
                    data[name] = m.group(1)
                    continue
                m = re.search(rf'<textarea name="{name}"[^>]*>\s*(.*?)</textarea>', html, re.S)
                if m:
                    data[name] = m.group(1)
            data["tab"] = key
            response = self.client.post(f"/abonents/5/edit/?tab={key}", data)
            self.assertEqual(response.status_code, 302, key)
        self.a.refresh_from_db()
        after = {f.attname: getattr(self.a, f.attname) for f in Potreb._meta.concrete_fields}
        self.assertEqual(before, after)

    def test_new_days_of_recalculation_reset_coefficient(self):
        self.login("Оператор")
        html = self.client.get("/abonents/5/edit/?tab=otop").content.decode()
        self.assertIn('name="dn_pereras"', html)
        data = {"tab": "otop", "tip_rasot": "2", "dn_pereras": "3", "tkf_perer": "1.100", "gkl_perer": "3.300",
                "nom_dog": "6", "name_potre": "ООО Крипс", "nom_dogab": "", "glav_potre": "Магазин"}
        self.client.post("/abonents/5/edit/", data)
        self.a.refresh_from_db()
        self.assertEqual(self.a.tip_rasot, 2)
        self.assertEqual(self.a.dn_pereras, 3)
        self.assertEqual(self.a.tkf_perer, 0)
        self.assertEqual(self.a.gkl_perer, 0)

    def test_accept_zero_is_not_allowed(self):
        self.login("Оператор")
        Potreb.objects.filter(potreb_id=9).update(accept=0)
        html = self.client.get("/abonents/9/edit/?tab=ishod").content.decode()
        accept = re.search(r'<select name="accept".*?</select>', html, re.S).group(0)
        self.assertNotIn('<option value="0"', accept)
        self.assertNotIn('<option value=""', re.search(r'<select name="accept".*?</select>', html, re.S).group(0))
        from abonents.forms import karta_form
        form = karta_form("ishod", {"accept": "0"}, instance=Potreb.objects.get(potreb_id=9))
        self.assertFalse(form.is_valid())
        self.assertIn("accept", form.errors)

    def test_topbar_shows_month_instead_of_search(self):
        self.login("Наблюдатель")
        html = self.client.get("/abonents/").content.decode()
        self.assertIn("<b>Октябрь 2026</b>", html)
        self.assertNotIn('class="gsearch"', html)
        self.assertIn("Система расчётов", html)

    def test_only_active_checkbox(self):
        Potreb.objects.filter(potreb_id=5).update(d_stat=60)
        Potreb.objects.filter(potreb_id=7).update(d_stat=63)
        self.login("Наблюдатель")
        # По умолчанию галочка включена.
        r = self.client.get("/abonents/?gorod=1")
        self.assertEqual(r.context["page"].paginator.count, 1)
        self.assertEqual(self.client.get("/abonents/?gorod=1&aktiv=0&aktiv=1").context["page"].paginator.count, 1)
        # Галочка запоминается.
        self.assertEqual(self.client.get("/abonents/?gorod=1").context["page"].paginator.count, 1)
        self.assertEqual(self.client.get("/abonents/?gorod=1&aktiv=0").context["page"].paginator.count, 3)

    def test_search(self):
        self.login("Наблюдатель")
        self.assertRedirects(self.client.get("/abonents/", {"gorod": 1, "dog": 8, "aktiv": 0}), "/abonents/9/")
        response = self.client.get("/abonents/", {"gorod": 1, "dog": 6})
        self.assertContains(response, "Магазин")
        self.assertContains(response, "Склад")
        response = self.client.get("/abonents/", {"gorod": 1, "q": "школа №1"})
        self.assertContains(response, "Школа")
        self.assertNotContains(response, "Крипс")
        self.assertRedirects(self.client.get("/abonents/", {"ls": 11}), "/abonents/11/")
        self.assertContains(self.client.get("/abonents/", {"ls": 99}, follow=True), "не найден")
        response = self.client.get("/abonents/", {"gorod": 1, "sub": 6})
        self.assertContains(response, "Склад")
        self.assertNotContains(response, "Магазин")

    def test_dogovor_card_totals(self):
        self.login("Наблюдатель")
        response = self.client.get("/abonents/dogovor/1/6/")
        self.assertContains(response, "Магазин")
        self.assertContains(response, "Склад")
        self.assertContains(response, "12,500")

    def test_only_admin_deletes(self):
        self.login("Оператор")
        self.assertEqual(self.client.post("/abonents/9/delete/").status_code, 403)
        self.client.logout()
        self.login("Администратор")
        self.assertRedirects(self.client.post("/abonents/9/delete/"), "/abonents/")
        self.assertFalse(Potreb.objects.filter(potreb_id=9).exists())


class NewAbonentAndPoteriTests(KartaBase):
    """Ввод нового абонента, тепловые потери и теплосчетчик."""

    def test_new_abonent_gets_next_number_and_defaults(self):
        self.login("Оператор")
        self.assertContains(self.client.get("/abonents/"), "Новый абонент")
        response = self.client.post("/abonents/new/", {
            "cod_gor": "1", "nom_dog": "77", "name_potre": "ООО Новое", "id_kategor": "1", "id_koteln": "16",
            "cod_minist": "1",
        })
        p = Potreb.objects.get(nom_dog=77)
        self.assertRedirects(response, f"/abonents/{p.potreb_id}/edit/?tab=ishod")
        self.assertEqual(p.potreb_id, 12)
        self.assertEqual(p.cod_gor_id, 1)
        self.assertEqual(p.accept, 1)
        self.assertEqual(p.name_kot, "(Р) Котельная №2")
        self.assertEqual(p.tip_rasot, 0)
        self.assertEqual(p.adres_sd, "")

    def test_new_abonent_needs_contract_and_city_kotel(self):
        self.login("Оператор")
        response = self.client.post("/abonents/new/", {"cod_gor": "2", "nom_dog": "", "id_koteln": "16"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Котельная из другого города")
        self.assertFalse(Potreb.objects.filter(potreb_id=12).exists())
        self.client.logout()
        self.login("Наблюдатель")
        self.assertEqual(self.client.get("/abonents/new/").status_code, 403)

    def test_poteri_raschet(self):
        from catalogs.models import SpDiametr, SpTeplpoteri
        SpDiametr.objects.create(id_diam=2, diametr=57)
        Potreb.objects.filter(potreb_id=5).update(dnotop=Decimal("30"), tip_rasot=1, nalich_pot=1)
        self.login("Оператор")
        self.assertContains(self.client.get("/abonents/5/"), "/abonents/5/poteri/")
        self.assertFalse(SpTeplpoteri.objects.filter(ln=5).exists())  # просмотр ничего не создаёт
        data = {"diametr1": "57", "dlina1": "100", "raschet": "1"}
        for i in range(2, 7):
            data.update({f"diametr{i}": "", f"dlina{i}": ""})
        self.assertRedirects(self.client.post("/abonents/5/poteri/", data), "/abonents/5/poteri/")
        sp = SpTeplpoteri.objects.get(ln=5)
        # 56 × 100 × 24 × 30 × 0,000001 = 4,032 → 4,03
        self.assertEqual(sp.poteri_pr1, Decimal("4.03"))
        self.assertEqual(sp.vs_poteri, Decimal("4.03"))
        self.a.refresh_from_db()
        self.assertEqual(self.a.poter_pr, Decimal("4.03"))
        self.assertEqual(self.a.poter_nag, 0)

    def test_teploschetchik_created_with_object_name(self):
        from catalogs.models import SpTeplosch
        self.login("Оператор")
        self.assertContains(self.client.get("/abonents/5/teploschetchik/"), 'value="ООО &quot;Крипс&quot;"')
        self.client.post("/abonents/5/teploschetchik/", {
            "name_obek": 'ООО "Крипс"', "adres_obek": "", "tip_prib": "СВТУ 10М", "kdata_pov": "2027-01-31",
        })
        obj = SpTeplosch.objects.get(ln=5)
        self.assertEqual(obj.tip_prib, "СВТУ 10М")
        self.assertEqual(str(obj.kdata_pov), "2027-01-31")
        self.client.logout()
        self.login("Наблюдатель")
        self.assertEqual(self.client.post("/abonents/5/teploschetchik/", {"tip_prib": "x"}).status_code, 403)

    def test_changing_kotel_updates_its_name(self):
        GvsKot.objects.create(id_kot=17, name_kot="(Р) Котельная №3", kod_gor=self.ryb)
        from abonents.forms import karta_form
        p = Potreb.objects.get(potreb_id=5)
        form = karta_form("ishod", None, instance=p)
        data = {name: form[name].value() for name in form.fields}
        data = {k: ("" if v is None else v) for k, v in data.items()}
        data["id_koteln"] = "17"
        form = karta_form("ishod", data, instance=p)
        self.assertTrue(form.is_valid(), form.errors)
        form.save()
        p.refresh_from_db()
        self.assertEqual(p.name_kot, "(Р) Котельная №3")


class RaschetTests(TestCase):
    """Расчет начислений на небольшом примере, посчитанном вручную."""

    def setUp(self):
        call_command("init_roles", stdout=open("/dev/null", "w"))
        self.ryb = Danrasch.objects.create(cod_gorod=1, name_gorod="Рыбница", tkoef=Decimal("0.5"),
                                           qpfakt=Decimal("46"), t_vozd=Decimal("10"), t_proekt=-17)
        self.dub = Danrasch.objects.create(cod_gorod=2, name_gorod="Дубоссары", qpfakt=Decimal("0"))
        common = dict(nalich_pot=3, dnotop=Decimal("30"), accept=1, id_katabon=1, d_stat=60)
        # По прибору: (20 − 10) × 1,5 = 15 Гкал, ГВС 2 и потери ГВС 1, лимит 40 при выставленных 10.
        self.a = Potreb.objects.create(potreb_id=1, nom_dog=1, cod_gor=self.ryb, tip_rasot=1, pok_old=10,
                                       pok_new=20, koef_pok=Decimal("1.5"), nach_gvs=2, poter_gvs=1, vd_limit=40,
                                       nach_finan=10, name_potre="По прибору", **common)
        # По нагрузке: 0,1 × 24 × 30 × 0,5 = 36 Гкал.
        self.b = Potreb.objects.create(potreb_id=2, nom_dog=2, cod_gor=self.ryb, tip_rasot=3,
                                       q_otoppr=Decimal("0.1"), **common)
        # По прибору и нагрузке: 4 Гкал по прибору и 0,05 × 24 × 30 × 0,5 = 18 по нагрузке.
        self.c = Potreb.objects.create(potreb_id=3, nom_dog=3, cod_gor=self.ryb, tip_rasot=2, pok_new=4,
                                       koef_pok=1, q_otoppr=Decimal("0.05"), **common)
        # Договор 50 в Дубоссарах: общий прибор (10 Гкал) делится по нагрузке 0,3 и 0,1.
        self.h = Potreb.objects.create(potreb_id=4, nom_dog=50, nom_dogab=50, nom_pribor=1, cod_gor=self.dub,
                                       tip_rasot=1, pok_new=10, koef_pok=1, q_otop=Decimal("0.3"), **common)
        self.s = Potreb.objects.create(potreb_id=5, nom_dog=51, nom_dogab=50, nom_pribor=1, cod_gor=self.dub,
                                       tip_rasot=1, q_otop=Decimal("0.1"), **common)
        # Договор расторгнут: расчёт его не трогает.
        self.x = Potreb.objects.create(potreb_id=6, nom_dog=6, cod_gor=self.ryb, tip_rasot=3,
                                       q_otoppr=Decimal("0.1"), nac_vsego=5, **{**common, "d_stat": 63})
        # ГВС по прибору: (1100 − 1000) × 0,0533 = 5,33 Гкал.
        self.g = Potreb.objects.create(potreb_id=7, nom_dog=7, cod_gor=self.dub, tip_rasot=3, tip_rasgvs=1,
                                       pok_newg=1100, pok_oldg=1000, **common)

    def login(self, role):
        user = get_user_model().objects.create_user(role, password="x")
        user.groups.add(Group.objects.get(name=role))
        self.client.force_login(user)

    def test_raschet(self):
        self.login("Оператор")
        response = self.client.post("/abonents/raschet/")
        self.assertRedirects(response, "/abonents/raschet/itog/")
        a, b, c, h, s = (Potreb.objects.get(pk=x.pk) for x in (self.a, self.b, self.c, self.h, self.s))
        self.assertEqual((a.nach_otpr, a.nach_otnag), (Decimal("15.000"), Decimal("0.000")))
        self.assertEqual((b.nach_otpr, b.nach_otnag), (Decimal("0.000"), Decimal("36.000")))
        self.assertEqual((c.nach_otpr, c.nach_otnag), (Decimal("4.000"), Decimal("18.000")))
        # По факту: 46 − 15 − 4 − 2 (ГВС) = 25 Гкал делятся пропорционально 36 и 18.
        self.assertEqual((b.nac_pogaz, c.nac_pogaz), (Decimal("16.667"), Decimal("8.333")))
        self.assertEqual(Danrasch.objects.get(cod_gorod=1).qpras, Decimal("54.00"))
        self.assertEqual((a.nach_otopl, b.nach_otopl, c.nach_otopl),
                         (Decimal("15.000"), Decimal("16.670"), Decimal("12.330")))
        self.assertEqual(a.nac_vsego, Decimal("18.000"))  # 15 + ГВС 2 + потери ГВС 1
        self.assertEqual((h.nach_otsub, s.nach_otsub, h.q_otopobsh), (Decimal("7.500"), Decimal("2.500"),
                                                                      Decimal("0.40000")))
        self.assertEqual((h.nach_otopl, s.nach_otopl), (Decimal("7.500"), Decimal("2.500")))
        # 10 + 18 × 2 ≥ 40, но 10 + 18 ≤ 40.
        self.assertEqual(list(Limit.objects.values_list("potreb_id", "koment")), [(1, "Лимит заканчивается")])
        x = Potreb.objects.get(pk=self.x.pk)
        # Расторгнутый счет не рассчитывается, но «Предъявлено к оплате» у него обнуляется:
        # иначе там остается итог договора от прежнего расчета.
        self.assertEqual((x.dnotop, x.nach_otnag, x.nac_vsego), (Decimal("30.0"), None, Decimal("0.000")))
        g = Potreb.objects.get(pk=self.g.pk)
        self.assertEqual((g.nach_gvs, g.nac_vsego), (Decimal("5.330"), Decimal("5.330")))
        self.assertEqual(SpTeplpoteri.objects.count(), 6)  # записи потерь создаются, как в старой программе
        page = self.client.get("/abonents/raschet/itog/")
        self.assertContains(page, "рассчитано абонентов 6")
        self.assertContains(page, "Без заключённого договора (не рассчитывались): 1")
        self.assertContains(page, "Лимит заканчивается")
        # Сумма у счета без договора (например, после загрузки из DBF) в итоги не входит,
        # но показана отдельно.
        self.assertNotContains(page, "Остались суммы «Предъявлено к оплате»")
        Potreb.objects.filter(pk=self.x.pk).update(nac_vsego=Decimal("5"))
        self.assertContains(self.client.get("/abonents/raschet/itog/"), "Остались суммы «Предъявлено к оплате»")
        itog = {label: values for label, values in page.context["table"]}
        ryb_vsego, url = itog["Предъявлено к оплате"][0]
        self.assertIn("pole=nac_vsego", url)
        rep = self.client.get(url).context["report"]
        self.assertEqual(rep.tables[0].total[-1], ryb_vsego)
        self.assertNotIn(self.x.potreb_id, [r[0] for r in rep.tables[0].sections[0].rows])
        self.assertEqual(self.client.get("/abonents/raschet/itog/rasshifr/?gorod=1&pole=zzz").status_code, 404)

    def test_viewer_cannot_run(self):
        self.login("Наблюдатель")
        self.assertEqual(self.client.get("/abonents/raschet/").status_code, 200)
        self.assertEqual(self.client.post("/abonents/raschet/").status_code, 403)

    def test_dan_rasch_sets_koef_and_days(self):
        self.login("Оператор")
        response = self.client.post("/abonents/raschet/dannye/?gorod=1", {
            "t_vozd": "4", "t_proekt": "-17", "dnotop": "31", "dngvs": "", "qpfakt": "46", "dnotop_vsem": "on",
        })
        self.assertRedirects(response, "/abonents/raschet/dannye/")
        d = Danrasch.objects.get(cod_gorod=1)
        self.assertEqual(d.tkoef, Decimal("0.400"))  # (18 − 4) / (18 + 17)
        self.assertEqual(set(Potreb.objects.filter(cod_gor=1, d_stat=60).values_list("dnotop", flat=True)),
                         {Decimal("31.0")})
        self.assertEqual(Potreb.objects.get(potreb_id=6).dnotop, Decimal("0.0"))
        self.assertEqual(set(Potreb.objects.filter(cod_gor=2).values_list("dnotop", flat=True)), {Decimal("30.0")})


class StatusTests(KartaBase):
    def test_status_change_and_gray_rows(self):
        self.login("Оператор")
        html = self.client.get("/abonents/9/edit/?tab=ishod").content.decode()
        self.assertIn('name="d_stat"', html)
        data = {name: value for name, value in re.findall(
            r'<(?:input|select|textarea)[^>]*name="([a-z_0-9]+)"[^>]*?(?:value="([^"]*)")?', html)}
        data.update({"tab": "ishod", "d_stat": "63", "accept": "3", "nom_dog": "8"})
        data.pop("csrfmiddlewaretoken", None)
        self.client.post("/abonents/9/edit/?tab=ishod", data)
        p = Potreb.objects.get(potreb_id=9)
        self.assertEqual((p.d_stat, p.d_stat_val), (63, "Расторгнут"))
        # Как в карте программы 2026 года: дата окончания и «отключен с» — сегодня.
        self.assertEqual((p.dk_dog, p.data_otkot), (date.today(), date.today()))
        self.assertContains(self.client.get("/abonents/9/"), "Лицевой счет: Расторгнут")
        listing = self.client.get("/abonents/?gorod=1&aktiv=0").content.decode()
        self.assertEqual(listing.count('class="clickable otkl"'), 1)


class PrisoedTests(KartaBase):
    def test_block_only_when_objects_exist(self):
        self.login("Наблюдатель")
        self.assertContains(self.client.get("/abonents/5/"), "Присоединённые объекты")  # posle_prib = 7
        page = self.client.get("/abonents/9/").content.decode()
        self.assertNotIn("<h2>Присоединённые объекты</h2>", page)
        self.assertNotIn("Присоединённый объект</summary>", page)  # наблюдателю добавлять нельзя
        self.assertEqual(self.client.post("/abonents/9/prisoed/", {"kind": "pribor", "ls": "5"}).status_code, 403)

    def test_add_into_free_slot_and_remove(self):
        self.login("Оператор")
        self.assertContains(self.client.get("/abonents/9/"), "Присоединённый объект</summary>")
        self.client.post("/abonents/9/prisoed/", {"kind": "nagruzka", "ls": "5"})
        self.client.post("/abonents/9/prisoed/", {"kind": "nagruzka", "ls": "7"})
        response = self.client.post("/abonents/9/prisoed/", {"kind": "nagruzka", "ls": "7"}, follow=True)
        self.assertContains(response, "уже присоединён")
        self.assertContains(self.client.post("/abonents/9/prisoed/", {"kind": "pribor", "ls": "999"}, follow=True),
                            "не найден")
        c = Potreb.objects.get(potreb_id=9)
        self.assertEqual((c.pris_nag1, c.pris_nag2, c.pris_nag3), (5, 7, None))
        self.assertContains(self.client.get("/abonents/9/"), 'href="/abonents/7/"')
        self.client.post("/abonents/9/prisoed/", {"remove": "pris_nag1"})
        c.refresh_from_db()
        self.assertEqual((c.pris_nag1, c.pris_nag2), (0, 7))
        # Освободившееся место занимается первым.
        self.client.post("/abonents/9/prisoed/", {"kind": "nagruzka", "ls": "11"})
        c.refresh_from_db()
        self.assertEqual(c.pris_nag1, 11)

    def test_sposob_switch(self):
        self.login("Наблюдатель")
        self.assertNotContains(self.client.get("/abonents/9/"), 'name="field" value="tip_rasot"')
        self.assertEqual(self.client.post("/abonents/9/sposob/", {"field": "tip_rasot", "value": "3"}).status_code, 403)
        self.login("Оператор")
        self.assertContains(self.client.get("/abonents/9/"), 'name="field" value="tip_rasot"')
        self.client.post("/abonents/9/sposob/", {"field": "tip_rasot", "value": "3"})
        self.assertEqual(Potreb.objects.get(potreb_id=9).tip_rasot, 3)
        self.client.post("/abonents/9/sposob/", {"field": "tip_rasot", "value": "7"})  # нет такого варианта
        self.assertEqual(Potreb.objects.get(potreb_id=9).tip_rasot, 3)
        response = self.client.post("/abonents/9/sposob/", {"field": "tip_rasgvs", "value": "1"})
        self.assertTrue(response["Location"].endswith("?tab=gvs"))
        self.assertEqual(Potreb.objects.get(potreb_id=9).tip_rasgvs, 1)

    def test_inactive_object_not_added(self):
        self.login("Оператор")
        Potreb.objects.filter(potreb_id=7).update(d_stat=61, d_stat_val="Приостановлен")
        response = self.client.post("/abonents/9/prisoed/", {"kind": "pribor", "ls": "7"}, follow=True)
        self.assertContains(response, "только действующий")
        c = Potreb.objects.get(potreb_id=9)
        self.assertFalse(c.posle_prib)


class NovMesTests(KartaBase):
    def tmp_path(self, name):
        import tempfile
        d = tempfile.mkdtemp()
        self.addCleanup(__import__("shutil").rmtree, d)
        return f"{d}/{name}"

    def setUp(self):
        super().setUp()
        from catalogs.models import Abot, Arhiv, Tarif
        self.Abot, self.Arhiv = Abot, Arhiv
        from abonents.models import TarifGorod
        Tarif.objects.create(id_tar=1, zn_tar=Decimal("412.50"))
        # Тариф берется у участка абонента (Рыбница), а не общий и не другого участка.
        TarifGorod.objects.create(gorod=1, id_tar=1, zn_tar=Decimal("420.00"))
        TarifGorod.objects.create(gorod=2, id_tar=1, zn_tar=Decimal("999.00"))
        Danrasch.objects.update(nom_mes=12, name_mes="Декабрь", god=2025)
        Potreb.objects.filter(potreb_id=5).update(nach_finan=Decimal("100"), nach_otopi=Decimal("40"),
                                                  vd_limit=Decimal("200"), ost_limit=Decimal("100"),
                                                  pok_old=Decimal("10"), dnotop=Decimal("31"),
                                                  pok_newg=Decimal("7"), pok_oldg=Decimal("3"))
        Potreb.objects.filter(potreb_id=7).update(nach_finan=Decimal("5"), nach_otopi=Decimal("5"))
        GvsKot.objects.update(nagr_kot=Decimal("1.5"))

    def test_close_month_and_undo(self):
        from abonents import nov_mes, reports
        Potreb.objects.filter(potreb_id=5).update(d_stat=60, nach_otopl=Decimal("3.25"), nach_gvs=Decimal("1"))
        before = {p.potreb_id: p for p in Potreb.objects.all()}
        polez = reports.polez_data()
        self.assertTrue(any(v for _, _, rows in polez for row in rows for v in row[1:]))
        r = nov_mes.run("тест")
        # В архив попадает полезный отпуск закрываемого месяца и тарифы.
        t = nov_mes.last().tables()
        self.assertEqual(set(t["polez_otpusk"]), {"", "gvs", "otopl"})
        self.assertEqual([[g, name, [[row[0]] + [Decimal(v) for v in row[1:]] for row in rows]]
                          for g, name, rows in t["polez_otpusk"][""]],
                         [[g, name, rows] for g, name, rows in polez])
        self.assertIn("tarif", t)
        self.assertIn("tarif_gorod", t)
        self.assertEqual((r.closed, r.opened), ((12, 2025), (1, 2026)))
        self.assertEqual(set(Danrasch.objects.values_list("nom_mes", "name_mes", "god")), {(1, "Январь", 2026)})
        a = Potreb.objects.get(potreb_id=5)
        # Январь: накопления за год и лимит обнуляются, затем прибавляется начисленное за декабрь.
        self.assertEqual((a.nach_finan, a.vd_limit, a.ost_limit), (Decimal("10.500"), 0, 0))
        self.assertEqual(a.nach_otopi, Decimal("50.500"))
        self.assertEqual((a.nac_vsego, a.dnotop, a.gkl_perer, a.tkf_perer), (0, 0, 0, 0))
        self.assertEqual((a.pok_old, a.pok_oldg), (Decimal("12.500"), Decimal("7")))
        self.assertEqual(GvsKot.objects.get().nagr_kot, 0)
        abot = self.Abot.objects.order_by("ras_kol")
        self.assertEqual([(x.ras_kol, x.koduslug, x.koduslug1, x.tarif, x.akcept, x.name_kateg, x.name_city) for x in abot],
                         [(Decimal("2.00"), 4, 4, Decimal("420.00"), 2, "Хозрасчет", "Рыбница"),
                          (Decimal("10.50"), 4, 4, Decimal("420.00"), 2, "Хозрасчет", "Рыбница")])
        self.assertEqual(list(self.Arhiv.objects.values_list("nom_mes", "god")), [(12, 2025)])

        nov_mes.undo()
        self.assertEqual(set(Danrasch.objects.values_list("nom_mes", "god")), {(12, 2025)})
        for p in Potreb.objects.all():
            for name in nov_mes.CHANGED:
                self.assertEqual(getattr(p, name), getattr(before[p.potreb_id], name), (p.potreb_id, name))
        self.assertEqual(GvsKot.objects.get().nagr_kot, Decimal("1.500"))
        self.assertFalse(self.Arhiv.objects.exists())
        self.assertFalse(nov_mes.MesArhiv.objects.exists())

    def test_ordinary_month_keeps_year_and_limit(self):
        from abonents import nov_mes
        Danrasch.objects.update(nom_mes=3, name_mes="Март", god=2026)
        nov_mes.run()
        a = Potreb.objects.get(potreb_id=5)
        self.assertEqual((a.nach_finan, a.nach_otopi, a.ost_limit), (Decimal("110.500"), Decimal("50.500"), Decimal("89.500")))
        Danrasch.objects.update(nom_mes=9, name_mes="Сентябрь", god=2026)
        nov_mes.run()  # открывается октябрь: накопление за сезон начинается заново
        a.refresh_from_db()
        self.assertEqual((a.nach_finan, a.nach_otopi), (Decimal("110.500"), 0))

    def test_page_and_permissions(self):
        self.login("Оператор")
        page = self.client.get("/abonents/nov-mes/")
        self.assertContains(page, "Декабрь 2025")
        self.assertContains(page, "Январь 2026")
        self.assertContains(page, "выполняет Администратор")
        self.assertEqual(self.client.post("/abonents/nov-mes/", {"ready": "1"}).status_code, 403)
        self.login("Администратор")
        self.assertContains(self.client.post("/abonents/nov-mes/", {}, follow=True), "Отметьте")
        response = self.client.post("/abonents/nov-mes/", {"ready": "1"}, follow=True)
        self.assertContains(response, "Выполнен переход к новому месяцу: Январь 2026")
        self.assertContains(response, "Отменить переход")
        from abonents import nov_mes
        import io
        from dbfread import DBF
        arh = nov_mes.last()
        dl = self.client.get(f"/abonents/nov-mes/{arh.pk}/1c/")
        self.assertIn('filename="Abpt2512.dbf"', dl["Content-Disposition"])
        path = self.tmp_path("Abpt2512.dbf")
        open(path, "wb").write(dl.content)
        recs = list(DBF(path, encoding="cp1251"))
        self.assertEqual(sorted(r["RAS_KOL"] for r in recs), [2.0, 10.5])
        self.assertEqual({(r["KODUSLUG"], r["NAME_KATEG"], r["CITY"]) for r in recs}, {(4, "Хозрасчет", 1)})
        response = self.client.post("/abonents/nov-mes/", {"action": "undo"}, follow=True)
        self.assertContains(response, "Переход отменён: снова Декабрь 2025")


class ReportsTests(KartaBase):
    def setUp(self):
        super().setUp()
        Potreb.objects.filter(potreb_id=5).update(nach_otopl=Decimal("8"), nach_gvs=Decimal("2.5"),
                                                  q_otoppr=Decimal("0.12"), pasp_got="2026-09-01")
        Potreb.objects.filter(potreb_id=9).update(q_otoppr=Decimal("0.3"), id_kategor=6)
        Potreb.objects.update(d_stat=60)

    def test_ot_kateg(self):
        from abonents import reports
        rep, _ = reports.build(reports.BY_KEY["ot_kateg"], self.ryb, {})
        t = rep.tables[0]
        # л/№ 7 — id_katabon 3 (присоединённый), л/№ 9 без начисления: не попадают.
        self.assertEqual([s.title for s in t.sections], ["ХОЗРАСЧЕТ"])
        self.assertEqual([r[1] for r in t.sections[0].rows], [5])
        self.assertEqual(t.sections[0].rows[0][4], "С")
        self.assertEqual(t.total[-1], Decimal("10.5"))
        self.assertEqual(rep.signatures, [f"Исполнитель {reports.LINE}",
                                          f"Ведущий специалист {reports.LINE} Борсуковский С.В."])
        # Подписанты берутся из справочника.
        from abonents.models import Podpisant
        Podpisant.objects.filter(rol="ved").update(fio="Петров П.П.")
        Podpisant.objects.create(rol="ved", gorod=2, doljnost="Инженер", fio="Сидоров С.С.")
        rep, _ = reports.build(reports.BY_KEY["ot_kateg"], self.ryb, {})
        self.assertTrue(rep.signatures[1].endswith("Петров П.П."))
        rep, _ = reports.build(reports.BY_KEY["ot_kateg"], self.dub, {})
        self.assertEqual(rep.signatures[1], f"Инженер {reports.LINE} Сидоров С.С.")

    def test_kotelnye_and_kategor(self):
        from abonents import reports
        rep, _ = reports.build(reports.BY_KEY["p_kot1"], self.ryb, {})
        row = rep.tables[0].sections[0].rows[0]
        self.assertEqual(row[0], "Котельная №2")
        self.assertEqual(row[-2:], [Decimal("2.5"), Decimal("8")])   # всего ГВС, отопление
        self.assertEqual(row[13:15], [Decimal("2.5"), Decimal("8")])  # хозрасчет
        rep, _ = reports.build(reports.BY_KEY["p_kot1"], self.ryb, {"vid": "gvs"})
        self.assertEqual(rep.tables[0].sections[0].rows[0][-2:], [Decimal("2.5"), Decimal("0")])
        Potreb.objects.filter(potreb_id=5).update(d_stat=63)   # без договора — не в отчете
        rep, _ = reports.build(reports.BY_KEY["p_kot1"], self.ryb, {})
        self.assertEqual(rep.tables[0].sections[0].rows, [])
        Potreb.objects.filter(potreb_id=5).update(d_stat=60)
        rep, _ = reports.build(reports.BY_KEY["p_kot2"], self.ryb, {})
        row = rep.tables[0].sections[0].rows[0]
        self.assertEqual(row[1], Decimal("0.3"))   # население абонентское — отдельно
        self.assertEqual(row[3], Decimal("0.12"))
        rep, _ = reports.build(reports.BY_KEY["potreb_kategor"], self.ryb, {})
        ryb = rep.tables[0].sections[0].rows[0]
        self.assertEqual((ryb[0], ryb[1], ryb[5], ryb[6]), ("г.Рыбница", Decimal("0.12"), Decimal("0.3"), Decimal("0.42")))
        rep, _ = reports.build(reports.BY_KEY["pasp_got"], self.ryb, {})
        self.assertEqual([r[0] for r in rep.tables[0].sections[0].rows], [5])
        self.assertEqual(rep.tables[1].total[1:], [1, 1])

    def test_polez_otpusk(self):
        from abonents import reports
        from catalogs.models import Tarif
        from abonents.models import TarifGorod
        Tarif.objects.create(id_tar=1, zn_tar=Decimal("985.76"), diz_tar=0)
        Tarif.objects.create(id_tar=2, zn_tar=Decimal("1"), diz_tar=0)
        TarifGorod.objects.create(gorod=1, id_tar=2, zn_tar=Decimal("691.32"))   # тариф участка важнее общего
        Kategorpot.objects.create(id_kotegor=2, name_kateg="Местный бюджет")
        Danrasch.objects.filter(cod_gorod=1).update(gkl_pod=Decimal("3"), gkl_bezpod=Decimal("1"))
        Potreb.objects.filter(potreb_id=5).update(poter_gvs=Decimal("0.5"))
        Potreb.objects.filter(potreb_id=9).update(nach_otopl=Decimal("4"), tip_rasot=2)
        Potreb.objects.create(potreb_id=21, name_potre="Школа", cod_gor=self.ryb, id_kategor=2, tip_rasot=1,
                              vid_topl=True, nach_otopl=Decimal("6"), d_stat=60)
        rep, _ = reports.build(reports.BY_KEY["polez_otpusk"], self.ryb, {})
        t = rep.tables[0]
        self.assertEqual(t.sections[0].rows[0][0], "Тариф (руб)")
        self.assertEqual(t.sections[0].rows[0][6], Decimal("985.76"))
        self.assertEqual(t.sections[0].rows[0][7:9], [Decimal("691.32"), Decimal("691.32")])
        pu, nagr = t.sections[1].rows
        self.assertEqual(pu[6], Decimal("11.0"))      # хозрасчет: 8 + 2,5 + потери ГВС 0,5
        self.assertEqual(pu[7], Decimal("6"))         # местный бюджет вместе с дизельным топливом
        self.assertIsNone(pu[-1])                     # «ВСЕГО» только в строке города
        self.assertEqual(t.sections[1].foot[-1], Decimal("21.0"))
        self.assertEqual(t.sections[1].foot[0], "Всего по г.Рыбница")
        self.assertEqual(nagr[2:4], [Decimal("3"), Decimal("1")])   # население абонентское 3:1
        itog = t.sections[-1]
        self.assertEqual([r[0] for r in itog.rows], ["П/У", "Нагр."])
        self.assertEqual((itog.foot[0], itog.foot[-1]), ("Итого по филиалу", Decimal("21.0")))
        # Шапка в два яруса, как в старом отчете, и так же в Excel.
        self.assertEqual(t.head[1][4], ("Местный", 1, 1))
        self.assertEqual(len(t.sections[0].rows[0]), len(reports.POLEZ_COLS))
        from io import BytesIO
        from openpyxl import load_workbook
        buf = BytesIO()
        reports.to_xlsx(rep, buf)
        ws = load_workbook(BytesIO(buf.getvalue())).active
        merged = {str(m) for m in ws.merged_cells.ranges}
        cell = next(c for row in ws.iter_rows() for c in row if c.value == "Бюджет (Гкал)")
        self.assertIn(f"{cell.coordinate}:{cell.offset(column=1).coordinate}", merged)
        self.assertNotIn("Местный, диз.", [c.value for row in ws.iter_rows() for c in row])
        # Расшифровка ячейки: местный бюджет, П/У Рыбницы — школа л/№ 21.
        url = t.sections[1].links[0][7]
        self.login("Наблюдатель")
        page = self.client.get(url)
        self.assertContains(page, "Школа")
        self.assertContains(page, 'href="/abonents/21/"')
        self.assertEqual(page.context["report"].tables[0].total[-1], Decimal("6"))
        # Итого по филиалу, ВСЕГО — все записи, сумма как в ячейке.
        rep_all = reports.polez_rasshifr({"uch": "0", "row": "all", "col": "11", "vid": ""})
        self.assertEqual(rep_all.tables[0].total[-1], Decimal("21.0"))
        self.assertEqual(self.client.get(url + "&xlsx=1").status_code, 200)
        self.assertEqual(self.client.get("/abonents/otchety/polez_otpusk/rasshifr/?col=99").status_code, 404)
        rep, _ = reports.build(reports.BY_KEY["polez_otpusk"], self.ryb, {"vid": "otopl"})
        self.assertEqual(rep.tables[0].sections[-1].foot[-1], Decimal("18"))

    def test_minist_and_pribory(self):
        from abonents import reports
        from catalogs.models import SetUrdevice, SetUrdevicelist
        Potreb.objects.filter(potreb_id=5).update(id_kategor=3, q_gvs=Decimal("0.01"))
        rep, _ = reports.build(reports.BY_KEY["nag_minist"], self.ryb, {})
        self.assertEqual([r[1] for r in rep.tables[0].sections[0].rows], ['ООО "Крипс"'])
        self.assertEqual(rep.tables[0].total[-1], Decimal("0.13"))
        Potreb.objects.filter(potreb_id=5).update(pok_old=Decimal("10"), pok_newg=Decimal("110"),
                                                  pok_oldg=Decimal("10"))
        SetUrdevicelist.objects.create(kod="3", name="ВКТ-7", koef=Decimal("2"), edn="Гкал")
        SetUrdevice.objects.create(p_kod="5", d_kod="3", d_num="123", d_pov=date(2022, 1, 1),
                                   d_next_pov=date(2026, 1, 1))
        rep, _ = reports.build(reports.BY_KEY["otch_pribor"], self.ryb, {})
        row = rep.tables[0].sections[0].rows[0]
        self.assertEqual((row[0], row[3], row[7]), (5, "ВКТ-7  №123", Decimal("5.0")))
        self.assertEqual(rep.signatures[0],
                         "Инспектор по контролю за потреблением тепловой энергии юридическим лицам")
        self.assertTrue(rep.signatures[1].endswith("Химач Н.И."))
        rep, _ = reports.build(reports.BY_KEY["otch_priborg"], self.ryb, {})
        self.assertEqual(rep.tables[0].sections[0].rows[0][6], Decimal("5.33"))
        rep, opts = reports.build(reports.BY_KEY["device_pov"], self.ryb, {"date": "2026-06-01"})
        self.assertEqual(opts["date"], "2026-06-01")
        self.assertEqual(rep.tables[0].sections[0].rows, [[5, 'ООО "Крипс"', date(2022, 1, 1), date(2026, 1, 1)]])
        rep, _ = reports.build(reports.BY_KEY["device_pov"], self.ryb, {"date": "2025-06-01"})
        self.assertEqual(rep.tables[0].sections[0].rows, [])

    def test_pages_and_excel(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from abonents import reports
        self.login("Наблюдатель")
        r = self.client.get("/abonents/otchety/")
        for d in reports.REPORTS:
            self.assertContains(r, f"/abonents/otchety/{d.key}/")
            page = self.client.get(f"/abonents/otchety/{d.key}/?gorod=1")
            self.assertEqual(page.status_code, 200, d.key)
        r = self.client.get("/abonents/otchety/ot_kateg/?gorod=1&kateg=1")
        self.assertContains(r, "ОТПУСК ТЕПЛОВОЙ ЭНЕРГИИ ДЛЯ АБОНЕНТОВ")
        self.assertContains(r, "10,500")
        x = self.client.get("/abonents/otchety/ot_kateg/?gorod=1&xlsx=1")
        self.assertEqual(x["Content-Disposition"], 'attachment; filename="ot_kateg_1.xlsx"')
        ws = load_workbook(BytesIO(x.content)).active
        values = [c for row in ws.iter_rows(values_only=True) for c in row if c is not None]
        self.assertIn("ИТОГО:", values)
        self.assertIn(10.5, values)
        self.assertEqual(self.client.get("/abonents/otchety/net/").status_code, 404)


class DokumentyTests(KartaBase):
    """«Информация-счёт» и акт по договору."""

    def setUp(self):
        super().setUp()
        from abonents.models import TarifGorod
        Kategorpot.objects.create(id_kotegor=4, name_kateg="Население бюджет")
        Danrasch.objects.filter(cod_gorod=1).update(nom_mes=10)
        TarifGorod.objects.create(gorod=1, id_tar=1, zn_tar=Decimal("412.37"), diz_tar=Decimal("900.00"))
        TarifGorod.objects.create(gorod=1, id_tar=3, zn_tar=Decimal("100.00"))
        TarifGorod.objects.create(gorod=1, id_tar=12, zn_tar=Decimal("50.00"))
        Potreb.objects.filter(potreb_id__in=[5, 7, 9]).update(d_stat=60)
        # Отопление с потерями 1.111, ГВС 0.5, потери ГВС не входят в сумму.
        Potreb.objects.filter(potreb_id=5).update(nach_otopl=Decimal("1.111"), nach_gvs=Decimal("0.5"),
                                                  poter_gvs=Decimal("9"), poterschob=Decimal("0.1"),
                                                  nameobakt="Магазин", adres_sd="Ленина", oba_dom="5")
        Potreb.objects.filter(potreb_id=7).update(nach_otopl=Decimal("2.003"), vid_topl=True)

    def test_summa_po_dogovoru(self):
        from abonents import documents
        dok = documents.sobrat(1, 6)
        self.assertEqual([s.p.potreb_id for s in dok.stroki], [5, 7])
        a, b = dok.stroki
        self.assertEqual((a.otopl, a.gvs, a.poteri, a.gkal), (Decimal("1.011"), Decimal("0.5"), Decimal("0.1"), Decimal("1.611")))
        self.assertEqual(a.tarif, Decimal("412.37"))
        self.assertEqual(a.itog, Decimal("1.611") * Decimal("412.37"))   # 664.32807
        self.assertEqual(a.itog_r, Decimal("664.33"))
        self.assertEqual(b.tarif, Decimal("900.00"))                      # дизельное топливо
        self.assertEqual(b.itog_r, Decimal("1802.70"))
        k = dok.kategorii[0]
        self.assertEqual(k.itog, Decimal("2467.03"))                      # 664.32807 + 1802.7
        self.assertEqual(k.otklonenie, Decimal("0.00"))
        self.assertEqual(dok.itog, Decimal("2467.03"))
        self.assertEqual(dok.mes_pr, "Октябре")
        self.assertIsNone(documents.sobrat(1, 999))

    def test_tarif_naselenie_s_vozmeshcheniem(self):
        from abonents import documents
        self.assertEqual(documents.tarif_kod(1, 4, 0), 3)
        self.assertEqual(documents.tarif_kod(1, 4, 1), 12)
        self.assertEqual(documents.tarif_kod(1, 1, 1), 1)
        self.assertEqual(documents.tarif_kod(2, 1, 0), 4)
        self.assertEqual(documents.tarif_kod(3, 7, 0), 10)

    def test_ne_dejstvuyushchie_ne_vhodyat(self):
        from abonents import documents
        Potreb.objects.filter(potreb_id=7).update(d_stat=63)
        self.assertEqual([s.p.potreb_id for s in documents.sobrat(1, 6).stroki], [5])
        self.assertEqual(documents.dogovory(1), [6])   # у договора 8 начисления нет

    def test_stranicy(self):
        self.login("Наблюдатель")
        r = self.client.get("/abonents/dogovor/1/6/schet/?ls=5")
        self.assertContains(r, "Информация-счёт")
        self.assertContains(r, "664,33")
        self.assertContains(r, "2467,03")
        self.assertContains(r, "Магазин ул.Ленина 5")
        r = self.client.get("/abonents/dogovor/1/6/akt/")
        self.assertContains(r, "в Октябре 2026 года было отпущено")
        self.assertEqual(self.client.get("/abonents/dogovor/1/6/xxx/").status_code, 404)
        self.assertEqual(self.client.get("/abonents/dogovor/1/999/akt/").status_code, 404)
        r = self.client.get("/abonents/dokumenty/?gorod=1&vid=akt")
        self.assertContains(r, "Договоров: 1")
        self.assertContains(self.client.get("/abonents/5/"), "/abonents/dogovor/1/6/schet/?ls=5")


class DokumentyPdfTests(DokumentyTests):
    def test_pdf(self):
        import tempfile
        from pathlib import Path
        from django.test import override_settings
        d = tempfile.mkdtemp()
        self.addCleanup(__import__("shutil").rmtree, d)
        self.login("Наблюдатель")
        with override_settings(TEPLO_DOC_DIR=d):
            r = self.client.get("/abonents/dogovor/1/6/schet/?ls=5&pdf=1")
            self.assertEqual(r["Content-Type"], "application/pdf")
            self.assertTrue(r.content.startswith(b"%PDF"))
            self.assertTrue((Path(d) / "ryb" / "2026" / "10" / "schet_6_5.pdf").exists())
            r = self.client.get("/abonents/dokumenty/?gorod=1&vid=akt&pdf=1")
            self.assertTrue(r.content.startswith(b"%PDF"))
            self.assertTrue((Path(d) / "ryb" / "2026" / "10" / "akt_6_5.pdf").exists())
        # Шрифт с кириллицей встроен (обычный и жирный), а не Helvetica.
        pdf_bytes = (Path(d) / "ryb" / "2026" / "10" / "akt_6_5.pdf").read_bytes()
        fonts = set(re.findall(rb"/BaseFont /[A-Z]{6}\+(\S+)", pdf_bytes))
        self.assertEqual(len(fonts), 2, fonts)
        self.assertTrue(any(b"Bold" in f for f in fonts), fonts)


class ArhivTests(KartaBase):
    """Отчеты за период по закрытым месяцам и загрузка архивов старой программы."""

    def setUp(self):
        super().setUp()
        from abonents import nov_mes
        Danrasch.objects.update(nom_mes=11, name_mes="Ноябрь", god=2026, t_vozd=Decimal("2.5"))
        Potreb.objects.filter(potreb_id__in=[5, 7, 9]).update(d_stat=60)
        Potreb.objects.filter(potreb_id=5).update(nach_otopl=Decimal("3"), nach_gvs=Decimal("1"), nac_vsego=Decimal("4"))
        nov_mes.run("тест")                       # закрыт ноябрь 2026
        Potreb.objects.filter(potreb_id=5).update(nach_otopl=Decimal("5"), nac_vsego=Decimal("5"))
        nov_mes.run("тест")                       # закрыт декабрь 2026

    def make_dbf(self, folder, name, spec, rows):
        import dbf
        t = dbf.Table(str(folder / name), spec, codepage="cp1251", dbf_type="vfp")
        t.open(dbf.READ_WRITE)
        for r in rows:
            t.append(r)
        t.close()

    def test_mesyacy(self):
        from abonents import arhiv
        self.assertEqual(arhiv.mesyacy((11, 2025), (2, 2026)), [(11, 2025), (12, 2025), (1, 2026), (2, 2026)])
        self.assertEqual(arhiv.poslednij(), (12, 2026))

    def test_ras_ab_po_mesyacam(self):
        rep, _ = reports.build(reports.BY_KEY["arh_ras_ab"], self.ryb, {"ls": "5", "s": "2026-10", "po": "2026-12"})
        rows = rep.tables[0].sections[0].rows
        self.assertEqual([r[0] for r in rows], ["ноябрь 2026", "декабрь 2026"])
        self.assertEqual([r[-1] for r in rows], [Decimal("4.000"), Decimal("5.000")])
        self.assertEqual(rows[0][1], Decimal("2.5"))
        self.assertEqual(rep.tables[0].total[-1], Decimal("9.000"))
        self.assertIn("октябрь 2026", rep.notes[0])

    def test_polez_za_period(self):
        rep, _ = reports.build(reports.BY_KEY["arh_polez"], self.ryb, {"s": "2026-11", "po": "2026-12"})
        t = rep.tables[0]
        self.assertEqual([s.title for s in t.sections], ["ноябрь 2026", "декабрь 2026"])
        # Хозрасчет (категория 1) — колонка «Всего хозрасчёт», ВСЕГО — последняя.
        self.assertEqual(t.sections[0].foot[-1], Decimal("4.000"))
        self.assertEqual(t.total[-1], Decimal("9.000"))

    def test_import_i_potreb_god(self):
        import tempfile
        from pathlib import Path
        from django.core.management import call_command
        from abonents import arhiv
        folder = Path(tempfile.mkdtemp())
        self.addCleanup(__import__("shutil").rmtree, folder)
        spec = ("potreb_id N(6,0); cod_gor N(2,0); nom_dog N(6,0); id_kategor N(2,0); id_katabon N(2,0); "
                "cod_minist N(3,0); name_potre C(60); nac_vsego N(12,3); nach_otopl N(12,3); tip_rasot N(1,0)")
        self.make_dbf(folder, "Po032024.dbf", spec, [(21, 1, 30, 2, 1, 1, "Школа № 5", 7.25, 7.25, 2),
                                                     (22, 1, 31, 2, 3, 1, "Не входит", 1, 1, 2)])
        self.make_dbf(folder, "Da032024.dbf", "cod_gorod N(2,0); name_gorod C(10); t_vozd N(5,1)", [(1, "Рыбница", -1.5)])
        self.make_dbf(folder, "po122026.dbf", spec, [(21, 1, 30, 2, 1, 1, "Школа № 5", 99, 99, 2)])
        out = io.StringIO()
        call_command("arhiv_import", dir=str(folder), stdout=out)
        self.assertIn("Загружено месяцев: 1, пропущено (уже есть): 1", out.getvalue())  # 12.2026 закрыт в вебе
        a = nov_mes_model().objects.get(nom_mes=3, god=2024)
        self.assertEqual((a.istochnik, a.potreb_count, a.nac_vsego), ("dbf", 2, Decimal("8.250")))
        from abonents import nov_mes
        self.assertEqual((nov_mes.last().nom_mes, nov_mes.last().god), (12, 2026))   # отменять можно только свой
        m = arhiv.zagruzit([(3, 2024)])[0][2]
        self.assertEqual(len(m.dejstvuyushchie(1)), 2)     # статуса договора в старом архиве нет — все записи
        rep, _ = reports.build(reports.BY_KEY["arh_potreb_god"], self.ryb, {"god": "2024"})
        rows = [r for s in rep.tables[0].sections for r in s.rows]
        self.assertEqual([(r[0], r[5], r[-1]) for r in rows], [("Школа № 5", Decimal("7.250"), Decimal("7.250"))])
        # Повторный запуск ничего не добавляет.
        call_command("arhiv_import", dir=str(folder), stdout=io.StringIO())
        self.assertEqual(nov_mes_model().objects.filter(nom_mes=3, god=2024).count(), 1)

    def test_excel(self):
        self.login("Наблюдатель")
        for key in ("arh_polez", "arh_ras_ab", "arh_potreb_god"):
            r = self.client.get(f"/abonents/otchety/{key}/?xlsx=1&ls=5&s=2026-11&po=2026-12&god=2026")
            self.assertEqual(r.status_code, 200, key)

    def test_stranica_otcheta(self):
        self.login("Наблюдатель")
        r = self.client.get("/abonents/otchety/arh_ras_ab/?ls=5&s=2026-11&po=2026-12")
        self.assertContains(r, 'type="month"')
        self.assertContains(r, "декабрь 2026")
        self.assertContains(self.client.get("/abonents/otchety/"), "Полезный отпуск за период")


def nov_mes_model():
    from abonents.models import MesArhiv
    return MesArhiv



class PochtaTests(DokumentyTests):
    """Отправка документов по почте и реестр писем."""

    def setUp(self):
        super().setUp()
        import tempfile
        from catalogs.models import SmtpNas
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(__import__("shutil").rmtree, self.tmp)
        SmtpNas.objects.create(smtp="smtp.example.org", port=465, from_dbf="teplo@example.org", tema="Информация",
                               pismo="Информация-счёт за", paunt=1, senduserna="teplo", password="x", puses=1)
        Potreb.objects.filter(potreb_id=5).update(mail="krips@example.org; buh@example.org", mail_ras=True)

    def test_spisok_i_rassylka(self):
        from django.core import mail
        from django.test import override_settings
        from abonents import pochta
        from abonents.models import Pismo
        self.login("Оператор")
        with override_settings(TEPLO_DOC_DIR=self.tmp, TEPLO_PISMO_NOMER=1001):
            rows = pochta.spisok(1, 10, 2026, "schet")
            self.assertEqual([(d.nom_dog, p.potreb_id, flag) for d, p, flag, z in rows], [(6, 5, True)])
            r = self.client.post("/abonents/rassylka/?gorod=1&vid=schet", {"dog": ["6"]}, follow=True)
            self.assertContains(r, "Отправлено: 1 из 1")
        self.assertEqual(len(mail.outbox), 1)
        m = mail.outbox[0]
        self.assertEqual(m.to, ["krips@example.org", "buh@example.org"])
        self.assertEqual(m.body, "Информация-счёт за Октябрь 2026 года")
        self.assertEqual(m.attachments[0][0], "schet_6_5.pdf")
        z = Pismo.objects.get()
        self.assertEqual((z.nomer, z.status, z.nom_mes, z.god), (1001, "ok", 10, 2026))
        self.assertIsNotNone(Potreb.objects.get(potreb_id=7).dataotpr)
        # Отправленному договору повторно не предлагается.
        self.assertFalse(pochta.spisok(1, 10, 2026, "schet")[0][2])
        r = self.client.get("/abonents/pisma/?mes=2026-10")
        self.assertContains(r, "1001")
        self.assertEqual(self.client.get(f"/abonents/pisma/{z.pk}/pdf/")["Content-Type"], "application/pdf")

    def test_oshibka_ne_preryvaet(self):
        from unittest import mock
        from django.test import override_settings
        from abonents import pochta
        from abonents.models import Pismo
        with override_settings(TEPLO_DOC_DIR=self.tmp), \
                mock.patch("django.core.mail.EmailMessage.send", side_effect=OSError("Connection refused")):
            res = pochta.rassylka(1, [6, 8], "schet", {})
        self.assertEqual([z.status for z in res], ["err", "err"])
        self.assertIn("Connection refused", res[0].oshibka)
        self.assertIn("Нет адреса", res[1].oshibka)        # у договора 8 нет почты
        self.assertEqual(Pismo.objects.count(), 2)
        self.assertTrue(pochta.spisok(1, 10, 2026, "schet")[0][2])   # неотправленный снова отмечен

    def test_nablyudatel_ne_otpravlyaet(self):
        self.login("Наблюдатель")
        self.assertEqual(self.client.post("/abonents/rassylka/?gorod=1", {"dog": ["6"]}).status_code, 403)
        self.assertEqual(self.client.post("/abonents/dogovor/1/6/schet/", {"action": "send"}).status_code, 403)


KREDO_XML = """<?xml version="1.0" encoding="windows-1251"?>
<NewDataSet>
  <xs:schema xmlns:xs="http://www.w3.org/2001/XMLSchema" id="NewDataSet">
    <xs:element name="Table"><xs:complexType><xs:sequence>
      <xs:element name="p_kod" type="xs:string"/><xs:element name="cur_pok" type="xs:decimal"/>
    </xs:sequence></xs:complexType></xs:element>
  </xs:schema>
  <Table><p_kod>5</p_kod><d_num>A-1</d_num><prev_pok>12.5</prev_pok><cur_pok>20,25</cur_pok><waitable>0</waitable><date_pok>2026-10-25T00:00:00+03:00</date_pok></Table>
  <Table><P_KOD>7 </P_KOD><D_NUM>B-2</D_NUM><PREV_POK>1</PREV_POK><CUR_POK>0.5</CUR_POK><DATE_POK>24.10.2026</DATE_POK></Table>
  <Table><p_kod>99999</p_kod><cur_pok>1</cur_pok></Table>
</NewDataSet>"""


class FakeInvoker:
    def __init__(self, otvet):
        self.otvet, self.calls = otvet, []

    def exec_proc(self, name, params):
        self.calls.append((name, params))
        return self.otvet


class KredoTests(KartaBase):
    """Приём показаний приборов из КРЭДО и разнос по лицевым."""

    def setUp(self):
        super().setUp()
        Danrasch.objects.filter(cod_gorod=1).update(nom_mes=10)
        Potreb.objects.filter(potreb_id=7).update(pok_new=Decimal("3"))

    def test_razobrat(self):
        from abonents import kredo
        rows = kredo.razobrat(KREDO_XML)
        self.assertEqual([r["p_kod"] for r in rows], ["5", "7", "99999"])
        self.assertEqual(rows[0]["cur_pok"], Decimal("20.25"))
        self.assertEqual(rows[0]["date_pok"], date(2026, 10, 25))
        self.assertEqual(rows[1]["date_pok"], date(2026, 10, 24))
        attr = kredo.razobrat('<root><row p_kod="5" cur_pok="13" date_pok="2026-10-01"/></root>')
        self.assertEqual(attr[0]["cur_pok"], Decimal("13"))
        with self.assertRaises(kredo.KredoError):
            kredo.razobrat("не xml")

    def test_zapros(self):
        from abonents import kredo
        inv = FakeInvoker("<r/>")
        kredo.zaprosit_pokazaniya(1, 2, 2028, inv)
        self.assertEqual(inv.calls, [("Kredo_DeviceMeters_Info",
                                      "@comorg=210,@org_id=null,@start_date=01.02.2028,@end_date=29.02.2028")])
        with self.assertRaises(kredo.KredoError) as e:
            kredo.zaprosit_pokazaniya(1, 10, 2026, FakeInvoker("<root><error>Нет прав</error></root>"))
        self.assertIn("Нет прав", str(e.exception))

    def test_sverka(self):
        from abonents import kredo
        st = {s.r["p_kod"]: s for s in kredo.sverit(kredo.razobrat(KREDO_XML))}
        self.assertEqual(st["5"].status, kredo.NOVOE)
        self.assertTrue(st["5"].otmetit)
        self.assertEqual(st["5"].rashod, Decimal("7.75"))
        self.assertEqual(st["7"].status, kredo.MENSHE)
        self.assertEqual(st["99999"].status, kredo.NET)

    def test_stranica_i_raznos(self):
        from unittest import mock
        from abonents import kredo
        from abonents.models import KredoPriem
        self.login("Оператор")
        with mock.patch.object(kredo, "ComInvoker", lambda cod_gor: FakeInvoker(KREDO_XML)):
            r = self.client.post("/abonents/kredo/?gorod=1", {"action": "kredo"}, follow=True)
        z = KredoPriem.objects.get()
        self.assertEqual((z.cod_gor, z.nom_mes, z.god, z.strok), (1, 10, 2026, 3))
        self.assertContains(r, "строк 3")
        self.assertContains(r, "меньше показания в карте")
        self.assertContains(r, "B-2")
        # Меньшее показание без подтверждения не разносится.
        r = self.client.post(f"/abonents/kredo/{z.pk}/", {"n": ["0", "1"]}, follow=True)
        self.assertContains(r, "подтвердите")
        self.assertEqual(Potreb.objects.get(potreb_id=5).pok_new, Decimal("12.5"))
        r = self.client.post(f"/abonents/kredo/{z.pk}/", {"n": ["0"]}, follow=True)
        self.assertContains(r, "разнесены: 1")
        a = Potreb.objects.get(potreb_id=5)
        self.assertEqual((a.pok_new, a.date_pok), (Decimal("20.25"), date(2026, 10, 25)))
        self.client.post(f"/abonents/kredo/{z.pk}/", {"n": ["1", "2"], "podtv": "1"})
        self.assertEqual(Potreb.objects.get(potreb_id=7).pok_new, Decimal("0.5"))
        z.refresh_from_db()
        self.assertEqual([i[0] for i in z.izmeneniya], [5, 7])
        self.assertIsNotNone(z.razneseno)

    def test_bez_pywin32_i_iz_faila(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from abonents.models import KredoPriem
        self.login("Оператор")
        r = self.client.post("/abonents/kredo/?gorod=1", {"action": "kredo"}, follow=True)
        self.assertContains(r, "КРЭДО:")
        self.assertFalse(KredoPriem.objects.exists())
        f = SimpleUploadedFile("otvet.xml", KREDO_XML.encode("cp1251"))
        r = self.client.post("/abonents/kredo/?gorod=1", {"action": "fail", "fail": f}, follow=True)
        self.assertEqual(KredoPriem.objects.get().istochnik, KredoPriem.FAIL)
        self.assertContains(r, "новое показание")

    def test_prosmotr_bez_prava(self):
        self.login("Наблюдатель")
        r = self.client.get("/abonents/kredo/?gorod=1")
        self.assertEqual(r.status_code, 200)
        self.assertNotContains(r, "Получить показания")
        self.assertEqual(self.client.post("/abonents/kredo/?gorod=1", {"action": "kredo"}).status_code, 403)


def _kredo_row(p_kod, d_num, prev, cur, data, doc):
    return (f"<row><org_id>AC23888B</org_id><comorg_id>210</comorg_id><liz>{p_kod}</liz><p_kod>{p_kod}</p_kod>"
            f"<d_kod>14</d_kod><d_tip>1</d_tip><d_num>{d_num}</d_num><tip_pok>2</tip_pok><prev_pok>{prev}</prev_pok>"
            f"<cur_pok>{cur}</cur_pok><date_pok>{data}</date_pok><waitable>0.0000</waitable><doc_id>{doc}</doc_id></row>")


# Как настоящий ответ Kredo_DeviceMeters_Info (31.10.2026): схема xsd, строки <row>,
# лицевой с ведущими нулями, одно показание повторяется по документам.
KREDO_REAL = ('<?xml version="1.0" encoding="windows-1251"?><records>'
              '<xsd:schema xmlns:xsd="http://www.w3.org/2001/XMLSchema" id="records">'
              '<xsd:element name="records"><xsd:complexType><xsd:choice maxOccurs="unbounded">'
              '<xsd:element name="Kredo_DeviceMeters_Info"><xsd:complexType><xsd:sequence>'
              '<xsd:element name="p_kod"/></xsd:sequence></xsd:complexType></xsd:element>'
              '</xsd:choice></xsd:complexType></xsd:element></xsd:schema>'
              + _kredo_row("000005", "997800", "12.5000", "13.0600", "2026-01-28", "A")
              + _kredo_row("000005", "997800", "12.5000", "13.0600", "2026-01-28", "B")
              + _kredo_row("000009", "2091363", "1.0", "6.3240", "2026-03-26", "C")
              + _kredo_row("000009", "2091363", "1.0", "4.1900", "2025-12-02", "D")
              + "</records>")


class KredoRealTests(KartaBase):
    def test_nastoyashchij_otvet(self):
        from abonents import kredo
        stroki = kredo.sverit(kredo.razobrat(KREDO_REAL))
        self.assertEqual([(s.potreb_id, s.r["cur_pok"]) for s in stroki],
                         [(5, Decimal("13.0600")), (9, Decimal("6.3240")), (9, Decimal("4.1900"))])
        self.assertEqual([s.otmetit for s in stroki], [True, True, False])
        self.assertFalse(stroki[0].dubl)


class NastroikaTests(KartaBase):
    """Папка для 1С в «Справочниках» → «Настройки программы»."""

    def test_papka_1c(self):
        from django.test import override_settings
        from abonents.models import Nastroika
        n = Nastroika.objects.get()
        n.papka_1c = r"\\server-ff\Abon_teplo1C"
        n.save()
        self.assertEqual(Nastroika.papka(), r"\\server-ff\Abon_teplo1C")
        n.papka_1c = ""
        n.save()
        with override_settings(TEPLO_1C_DIR=r"D:\1c"):
            self.assertEqual(Nastroika.papka(), r"D:\1c")

    def test_prava(self):
        from abonents.models import Nastroika
        pk = Nastroika.objects.get().pk
        self.login("Оператор")
        self.assertEqual(self.client.get("/catalogs/nastroika/").status_code, 200)
        self.assertEqual(self.client.get(f"/catalogs/nastroika/{pk}/edit/").status_code, 403)
        self.client.logout()
        self.login("Администратор")
        r = self.client.post(f"/catalogs/nastroika/{pk}/edit/", {"papka_1c": r"\\srv\1c"})
        self.assertIn(r.status_code, (200, 302))
        self.assertEqual(Nastroika.objects.get().papka_1c, r"\\srv\1c")


class DokumentyArhivTests(DokumentyTests):
    """Документы за закрытый месяц из архива."""

    def setUp(self):
        super().setUp()
        from abonents.models import MesArhiv, TarifGorod
        a = MesArhiv(nom_mes=9, god=2026, istochnik=MesArhiv.WEB)
        a.set_tables({"potreb": list(Potreb.objects.order_by("pk").values()),
                      "danrasch": list(Danrasch.objects.order_by("pk").values()),
                      "tarif": [], "tarif_gorod": list(TarifGorod.objects.values())})
        a.save()

    def izmenit(self):
        """После закрытия сентября изменились начисления и тариф."""
        from abonents.models import TarifGorod
        Potreb.objects.filter(potreb_id=5).update(nach_otopl=Decimal("3"))
        TarifGorod.objects.filter(gorod=1, id_tar=1).update(zn_tar=Decimal("500"))

    def test_sobrat_po_arhivu(self):
        from abonents import arhiv, documents
        self.izmenit()
        m = arhiv.najti((9, 2026))
        dok = documents.sobrat(1, 6, mesyac=m)
        self.assertEqual((dok.mes, dok.god, dok.mes_pr), (9, 2026, "Сентябре"))
        self.assertTrue(dok.arhiv)
        self.assertFalse(dok.tarif_tekushchij)
        self.assertEqual(dok.stroki[0].itog_r, Decimal("664.33"))      # архивные Гкал и тариф
        self.assertEqual(dok.itog, Decimal("2467.03"))
        self.assertEqual(documents.dogovory(1, mesyac=m), [6])
        self.assertEqual(arhiv.vybor(), [("2026-09", "сентябрь 2026")])
        self.assertNotEqual(documents.sobrat(1, 6).stroki[0].itog_r, Decimal("664.33"))

    def test_bez_tarifov_v_arhive(self):
        from abonents import arhiv, documents
        from abonents.models import MesArhiv
        a = MesArhiv.objects.get()
        t = a.tables()
        t["tarif_gorod"] = []
        a.set_tables(t)
        a.save()
        self.izmenit()
        dok = documents.sobrat(1, 6, mesyac=arhiv.najti((9, 2026)))
        self.assertTrue(dok.tarif_tekushchij)
        self.assertEqual(dok.stroki[0].tarif, Decimal("500"))

    def test_stranicy_i_pdf(self):
        import tempfile
        from pathlib import Path
        from django.test import override_settings
        self.izmenit()
        self.login("Наблюдатель")
        r = self.client.get("/abonents/dogovor/1/6/akt/?ls=5&mes=2026-09")
        self.assertContains(r, "в Сентябре 2026 года было отпущено")
        self.assertContains(r, "закрытый месяц")
        self.assertContains(r, "ls=5&amp;mes=2026-09&amp;pdf=1")
        self.assertEqual(self.client.get("/abonents/dogovor/1/6/akt/?mes=2025-01").status_code, 404)
        r = self.client.get("/abonents/dokumenty/?gorod=1&vid=schet&mes=2026-09")
        self.assertContains(r, "664,33")
        d = tempfile.mkdtemp()
        self.addCleanup(__import__("shutil").rmtree, d)
        with override_settings(TEPLO_DOC_DIR=d):
            r = self.client.get("/abonents/dogovor/1/6/schet/?mes=2026-09&pdf=1")
            self.assertTrue(r.content.startswith(b"%PDF"))
            self.assertTrue((Path(d) / "ryb" / "2026" / "9" / "schet_6_5.pdf").exists())

    def test_rassylka_za_arhivnyj_mesyac(self):
        from django.core import mail
        from django.test import override_settings
        from catalogs.models import SmtpNas
        from abonents.models import Pismo
        import tempfile
        d = tempfile.mkdtemp()
        self.addCleanup(__import__("shutil").rmtree, d)
        SmtpNas.objects.create(smtp="smtp.example.org", port=25, from_dbf="t@example.org", pismo="Счёт за")
        from abonents.models import MesArhiv
        a = MesArhiv.objects.get()
        t = a.tables()
        for row in t["potreb"]:
            if row["potreb_id"] == 5:
                row["mail"], row["mail_ras"] = "krips@example.org", True
        a.set_tables(t)
        a.save()
        self.login("Оператор")
        r = self.client.get("/abonents/rassylka/?gorod=1&vid=schet&mes=2026-09")
        self.assertContains(r, "сентябрь 2026")
        with override_settings(TEPLO_DOC_DIR=d):
            self.client.post("/abonents/rassylka/?gorod=1&vid=schet&mes=2026-09", {"dog": ["6"], "mes": "2026-09"})
        z = Pismo.objects.get()
        self.assertEqual((z.nom_mes, z.god, z.status), (9, 2026, Pismo.OK))
        self.assertIn("Сентябрь 2026", mail.outbox[0].body)
