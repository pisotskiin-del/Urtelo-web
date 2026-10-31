"""Отчеты старой программы: форма otchet (кнопки «Отчеты») и меню «Печать»
(pasp_got.prg, potreb_koteln.prg, potreb_kategor.prg). Каждый отчет —
функция, которая по городу и параметрам строит Report: заголовок, таблицы
с группами, итогами по группам и общим итогом, подписи. Один и тот же
Report показывается на экране, печатается и выгружается в Excel."""

import datetime
from dataclasses import dataclass, field
from decimal import Decimal

from django.db.models import Q
from django.utils import timezone

from catalogs.models import (Akcept, Danrasch, GvsKot, Kategorpot, Limit, Minist, Potreb, SetUrdevice,
                             SetUrdevicelist)

from .layout import RASCHET, TIP_RASOT

LINE = "__________________________________"


def _signer(rol, city=None):
    from .models import Podpisant
    return Podpisant.get(rol, city)


def _signatures(city=None):
    """Подписи «Исполнитель» и «Ведущий специалист» из справочника
    «Подписанты документов»."""
    from .models import Podpisant
    out = []
    for rol, default in ((Podpisant.ISPOLNITEL, "Исполнитель"), (Podpisant.VED_SPEC, "Ведущий специалист")):
        dolj, fio = _signer(rol, city)
        out.append(f"{dolj or default} {LINE} {fio}".strip())
    return out


FILIAL = 'ПЕРЕЧЕНЬ ПОТРЕБИТЕЛЕЙ МГУП "ТИРАСТЕПЛОЭНЕРГО" ФИЛИАЛ В Г.РЫБНИЦА'
PERECHEN = 'ПЕРЕЧЕНЬ ПОТРЕБИТЕЛЕЙ МГУП "ТИРАСТЕПЛОЭНЕРГО"'


@dataclass
class Col:
    title: str
    kind: str = "text"      # text, int, num, date
    places: int = 3
    total: bool = False     # суммировать в итогах


@dataclass
class Section:
    title: str = ""
    rows: list = field(default_factory=list)
    foot: list = None       # строка итога группы
    links: list = None      # ссылки на расшифровку ячеек: по строке список url или None
    foot_links: list = None
    mark_rows: bool = False  # строки выделены цветом и выше обычных (строка тарифов)
    mark_foot: bool = False  # так же выделенный итог (итог по филиалу)


@dataclass
class Table:
    cols: list
    sections: list
    groups: list = None     # верхняя строка шапки: [(заголовок, колонок)]
    total: list = None      # строка общего итога
    title: str = ""
    head: list = None       # многоярусная шапка: строки [(заголовок, колонок, строк)] вместо groups и cols


@dataclass
class Report:
    title: str
    lines: list
    tables: list
    notes: list = field(default_factory=list)
    signatures: list = field(default_factory=list)
    landscape: bool = False
    approve: list = field(default_factory=list)   # гриф «УТВЕРЖДАЮ» справа вверху


@dataclass
class ReportDef:
    key: str
    title: str
    group: str
    build: object
    options: list = field(default_factory=list)   # [(имя, подпись, [(значение, подпись)] | None | DATE)]
    per_city: bool = True
    hint: str = ""


def n(value):
    return value if value is not None else Decimal(0)


def _sum_row(cols, rows, label, label_col=None):
    """Строка итога: подпись в первой текстовой колонке (или label_col),
    суммы в колонках с total=True."""
    if label_col is None:
        label_col = next((i for i, c in enumerate(cols) if c.kind == "text"), 0)
    out = [None] * len(cols)
    for i, c in enumerate(cols):
        if c.total:
            out[i] = sum((n(r[i]) for r in rows), Decimal(0))
    out[label_col] = label
    return out


def _numbered(rows):
    return [[i] + r for i, r in enumerate(rows, 1)]


def _city_lines(city, month=True):
    lines = [f"г. {(city.name_gorod or '').strip()}"]
    if month:
        lines.append(f"за {(city.name_mes or '').strip()} {city.god} г.")
    return lines


def _active(qs):
    """Только лицевые счета с заключённым договором (программа 2026 года
    отбирает d_stat=60; возобновлённые считаются вместе с ними, как в расчете)."""
    return qs.filter(d_stat__in=RASCHET)


def _approve(year=True):
    lines = ["УТВЕРЖДАЮ:"]
    if year:
        lines.append(f'"____" ______________ {datetime.date.today().year} г.')
    return lines


def _kategor():
    return dict(Kategorpot.objects.values_list("id_kotegor", "name_kateg"))


def _akcept():
    return dict(Akcept.objects.values_list("id_akc", "name_akc"))


def _grouped(items, key, title, build_row, cols):
    """Разбивает записи на группы по key (порядок — как в items) и
    добавляет итог каждой группы."""
    sections, current, keyval = [], None, object()
    for item in items:
        k = key(item)
        if k != keyval:
            current = Section(title(k))
            sections.append(current)
            keyval = k
        current.rows.append(build_row(item))
    for s in sections:
        for i, r in enumerate(s.rows, 1):
            r[0] = i
        s.foot = _sum_row(cols, s.rows, f"всего: {s.title}", label_col=3)
    return sections


# --- Отчеты формы «Отчеты» ---------------------------------------------------

def ot_kateg(city, opts):
    """Кнопка «Распределение по бюджетам» (отчет OT_KATEG): начислено за месяц
    по категориям потребителей. Ветка с предоплатой (pr_xozr) не переносится:
    предоплата не используется."""
    qs = Potreb.objects.filter(cod_gor=city.cod_gorod).exclude(id_katabon=3).exclude(nac_vsego=0).exclude(nac_vsego=None)
    kat = opts.get("kateg")
    if kat:
        qs = qs.filter(id_kategor=int(kat))
    else:
        qs = qs.filter(nom_dog__gt=0)
    names, akc = _kategor(), _akcept()
    cols = [Col("№ п/п", "int"), Col("Лиц. №", "int"), Col("Код бух.", "int"), Col("Абонент"), Col("Акцепт"),
            Col("Выстав. аванс", "num", total=True), Col("Отпуск (Гкал)", "num", total=True),
            Col("К оплате (Гкал)", "num", total=True)]
    sections = _grouped(
        qs.order_by("id_kategor", "potreb_id"), lambda p: p.id_kategor,
        lambda k: (names.get(k) or f"Категория {k}").upper(),
        lambda p: [0, p.potreb_id, p.cod_buhg, p.name_potre, akc.get(p.accept), n(p.predopl),
                   n(p.nac_vsego) + n(p.predopl), n(p.nac_vsego)],
        cols)
    rows = [r for s in sections for r in s.rows]
    return Report("ОТПУСК ТЕПЛОВОЙ ЭНЕРГИИ ДЛЯ АБОНЕНТОВ", _city_lines(city),
                  [Table(cols, sections, total=_sum_row(cols, rows, "ИТОГО:", label_col=4))],
                  signatures=_signatures(city))


def ot_minist(city, opts):
    """Кнопка «Расшифровка по министерствам»: республиканский бюджет по
    министерствам."""
    qs = Potreb.objects.filter(cod_gor=city.cod_gorod, nom_dog__gt=0, id_kategor=3).exclude(id_katabon=3)
    names, akc = dict(Minist.objects.values_list("cod_minist", "name_minis")), _akcept()
    cols = [Col("№ п/п", "int"), Col("Код АС", "int"), Col("Код бух.", "int"), Col("Абонент"), Col("Акцепт"),
            Col("Выстав. аванс", "num", total=True), Col("Отпуск (Гкал)", "num", total=True),
            Col("К оплате (Гкал)", "num", total=True)]
    sections = _grouped(
        qs.order_by("cod_minist", "potreb_id"), lambda p: p.cod_minist_id,
        lambda k: (names.get(k) or "Министерство не указано").strip().upper(),
        lambda p: [0, p.cod_ac, p.cod_buhg, p.name_potre, akc.get(p.accept), n(p.predopl),
                   n(p.nac_vsego), n(p.nac_vsego) - n(p.predopl)],
        cols)
    rows = [r for s in sections for r in s.rows]
    return Report("РЕСПУБЛИКАНСКИЙ БЮДЖЕТ", ["Расшифровка по министерствам"] + _city_lines(city),
                  [Table(cols, sections, total=_sum_row(cols, rows, "ИТОГО:", label_col=4))],
                  signatures=_signatures(city))


def ved_potreb(city, opts):
    """Кнопка «Ведомость по приборам»: абоненты с показаниями прибора, у
    которых есть (или нет) начисления по разнице показаний."""
    qs = Potreb.objects.filter(cod_gor=city.cod_gorod, pok_new__gt=0)
    nenach = opts.get("vid") == "nenach"
    qs = qs.filter(Q(nach_otpr__lte=0) | Q(nach_otpr=None)) if nenach else qs.filter(nach_otpr__gt=0)
    cols = [Col("№ догов.", "int"), Col("Наименование потребителя"), Col("разн. показаний", "num", 2, total=True)]
    rows = [[p.nom_dog, p.name_potre, n(p.nach_otpr)] for p in qs.order_by("nom_dog", "potreb_id")]
    return Report("Ведомость начисления по приборам", _city_lines(city) + ["не начислено" if nenach else "начислено"],
                  [Table(cols, [Section(rows=rows)], total=_sum_row(cols, rows, "Итого:"))])


def nach_nul(city, opts):
    """Кнопка «Без начисления»: абоненты, у которых нет начисления за
    отопление (или за ГВС)."""
    name = "nach_gvs" if opts.get("vid") == "gvs" else "nach_otopl"
    qs = Potreb.objects.filter(cod_gor=city.cod_gorod).filter(Q(**{f"{name}__lte": 0}) | Q(**{name: None}))
    cols = [Col("№ дог.", "int"), Col("Наименование потребителей"), Col("предоплата", "num"),
            Col("начисл. отопление", "num"), Col("Начислено ГВС", "num")]
    rows = [[p.nom_dog, p.name_potre, n(p.predopl), n(p.nach_otopl), n(p.nach_gvs)]
            for p in qs.order_by("nom_dog", "potreb_id")]
    return Report("Список потребителей без начисления",
                  _city_lines(city) + ["по ГВС" if name == "nach_gvs" else "по отоплению"],
                  [Table(cols, [Section(rows=rows)])], notes=[f"Всего: {len(rows)}"])


def _koteln():
    return {k.id_kot: k.name_kot for k in GvsKot.objects.exclude(id_kot=None)}


KOT_KATEG = [(4, "население ведомственное"), (6, "население абонентское"), (5, "население хозрасчет"),
             (2, "местный бюджет"), (3, "республиканский бюджет"), (7, "ОСКиИОД"), (1, "хозрасчет")]
VID_USLUGA = [("", "Отопление и ГВС"), ("gvs", "Без отопления (только ГВС)"), ("otopl", "Без ГВС (только отопление)")]


def _gvs(p):
    """ГВС в отчетах программы 2026 года — начислено плюс тепловые потери ГВС."""
    return n(p.nach_gvs) + n(p.poter_gvs)


def p_kot1(city, opts):
    """Кнопка «Потребление по котельным» (p_kot1n): начислено за ГВС (с
    потерями ГВС) и отопление по котельным и категориям. Котельная попадает
    в отчет, если у нее есть начисление по выбранной услуге."""
    vid = opts.get("vid")
    names = _koteln()
    data, bez_kot = {}, []
    for p in _active(Potreb.objects.filter(cod_gor=city.cod_gorod)).order_by("id_koteln", "potreb_id"):
        g = Decimal(0) if vid == "otopl" else _gvs(p)
        o = Decimal(0) if vid == "gvs" else n(p.nach_otopl)
        has = (n(p.nach_gvs) > 0 if vid == "gvs" else n(p.nach_otopl) > 0 if vid == "otopl"
               else n(p.nach_gvs) > 0 or n(p.nach_otopl) > 0)
        if not p.id_koteln:
            if n(p.nach_gvs) > 0 or n(p.nach_otopl) > 0:
                bez_kot.append(p.potreb_id)
            continue
        if p.id_koteln not in names:
            continue
        row = data.setdefault(p.id_koteln, {"has": False})
        row["has"] = row["has"] or has
        for k in (p.id_kategor, "all"):
            pg, po = row.get(k, (Decimal(0), Decimal(0)))
            row[k] = (pg + g, po + o)
    cols = [Col("Наименование котельной")]
    groups = [("", 1)]
    for _, title in KOT_KATEG + [(None, "Всего")]:
        cols += [Col("ГВС", "num", 2, total=True), Col("Отопл.", "num", 2, total=True)]
        groups.append((title, 2))
    rows = []
    for kot, row in data.items():
        if not row["has"]:
            continue
        r = [(names[kot] or "")[3:].strip()]
        for k, _ in KOT_KATEG + [("all", "")]:
            r += list(row.get(k, (Decimal(0), Decimal(0))))
        rows.append(r)
    total = _sum_row(cols, rows, "Итого")
    vsego = (total[-2] or 0) + (total[-1] or 0)
    notes = [f"ВСЕГО ПОТРЕБЛЕНИЕ ПО КОТЕЛЬНЫМ: {vsego:.2f} ГКал"]
    if bez_kot:
        notes.append("Не указан № котельной у потребителей: " + ", ".join(map(str, bez_kot)))
    lines = [" ".join(_city_lines(city))] + ([dict(VID_USLUGA)[vid]] if vid else [])
    return Report("ФАКТИЧЕСКОЕ ПОТРЕБЛЕНИЕ ПО КОТЕЛЬНЫМ", lines,
                  [Table(cols, [Section(rows=rows)], groups=groups, total=total)],
                  notes=notes, signatures=_signatures(city), landscape=True)


def p_kot2(city, opts):
    """Кнопка «Нагрузка по котельным» (p_kot2): часовая нагрузка по
    котельным, население (абонентское) отдельно от прочих."""
    names = _koteln()
    data = {}
    for p in _active(Potreb.objects.filter(cod_gor=city.cod_gorod)).order_by("id_koteln"):
        if p.id_koteln not in names:
            continue
        r = data.setdefault(p.id_koteln, [Decimal(0)] * 4)
        i = 0 if p.id_kategor == 6 else 2
        r[i] += n(p.q_otoppr)
        r[i + 1] += n(p.q_gvs)
    cols = [Col("Наименование котельной")] + [Col(t, "num", 4, total=True) for t in
                                               ("Отопл.", "ГВС", "Отопл.", "ГВС", "Отопл.", "ГВС", "ВСЕГО")]
    groups = [("", 1), ("Население", 2), ("Прочие", 2), ("", 3)]
    rows = [[names[k], a, b, c, d, a + c, b + d, a + b + c + d] for k, (a, b, c, d) in data.items()]
    total = _sum_row(cols, rows, "Итого")
    return Report("РАСПРЕДЕЛЕННАЯ НАГРУЗКА ПО КОТЕЛЬНЫМ", _city_lines(city),
                  [Table(cols, [Section(rows=rows)], groups=groups, total=total)],
                  notes=[f"ВСЕГО НАГРУЗКА ПО КОТЕЛЬНЫМ: {total[-1]:.4f} Гкал/час"],
                  signatures=_signatures(city))


def nag_kateg(city, opts):
    """Кнопка «Нагрузка по категориям» (nag_kateg): часовая нагрузка на
    отопление абонентов по категориям."""
    names = _kategor()
    cols = [Col("№ п/п", "int"), Col("Л/Н", "int"), Col("Абонент"), Col("Нагрузка, Гкал/час", "num", 5, total=True)]
    qs = Potreb.objects.filter(cod_gor=city.cod_gorod, q_otoppr__gt=0).order_by("id_kategor", "potreb_id")
    sections = []
    for s in _grouped(qs, lambda p: p.id_kategor, lambda k: (names.get(k) or f"Категория {k}").upper(),
                      lambda p: [0, p.potreb_id, p.name_potre, n(p.q_otoppr)], cols):
        s.foot = _sum_row(cols, s.rows, f"всего: {s.title}", label_col=2)
        sections.append(s)
    rows = [r for s in sections for r in s.rows]
    return Report("ТЕПЛОВАЯ НАГРУЗКА ПО КАТЕГОРИЯМ ПОТРЕБИТЕЛЕЙ", _city_lines(city),
                  [Table(cols, sections, total=_sum_row(cols, rows, "ИТОГО:", label_col=2))],
                  signatures=_signatures(city))


def limit(city, opts):
    """Кнопка «Превышение лимита»: таблица limit, её заполняет расчет."""
    cols = [Col("№ лиц.сч.", "int"), Col("Потребитель"), Col("Текущее начисление", "num"),
            Col("Установлен лимит", "num"), Col("Начислено с начала фин. года", "num"), Col("Коментарий")]
    rows = [[x.potreb_id, x.name_dog, n(x.nac_vsego), n(x.vd_limit), n(x.nach_finan), x.koment]
            for x in Limit.objects.order_by("potreb_id")]
    return Report("Список потребителей с заканчивающимся лимитом", [_today()],
                  [Table(cols, [Section(rows=rows)])])


# --- Меню «Печать» ------------------------------------------------------------

def _today():
    return datetime.date.today().strftime("%d.%m.%Y")


def _name_glav(p):
    name, glav = (p.name_potre or "").strip(), (p.glav_potre or "").strip()
    return f"{name}  ({glav})" if glav else name


def pasp_got(city, opts):
    """«Печать → Паспорта готовности» (pasp_got.prg): объекты хозрасчета и
    бюджета с отопительной нагрузкой и даты паспортов готовности. Колонки
    «Дата запуска» и «Количество дней отпуска» в старом отчете пустые — их
    заполняют от руки. В отчет входят только заключённые договоры."""
    qs = _active(Potreb.objects.filter(cod_gor=city.cod_gorod, id_kategor__lt=4, q_otoppr__gt=0)).order_by("potreb_id")
    cols = [Col("Лиц.№", "int"), Col("Наименование потребителя"), Col("Паспорт готовности", "date"),
            Col("Дата запуска"), Col("Количество дней отпуска"), Col("Адрес расположения объекта")]
    recs = list(qs)
    rows = [[p.potreb_id, _name_glav(p), p.pasp_got, "", "",
             f"ул. {(p.adres_sd or '').strip()}. дом {(p.oba_dom or '').strip()}. {(p.oba_kv or '').strip()}"]
            for p in recs]
    scols = [Col(""), Col("Количество объектов", "int", total=True), Col("С паспортами готовности", "int", total=True)]
    srows = [[title, sum(1 for p in recs if p.id_kategor == k), sum(1 for p in recs if p.id_kategor == k and p.pasp_got)]
             for k, title in ((1, "Хозрасчет"), (2, "Местный бюджет"), (3, "Республиканский бюджет"))]
    return Report(PERECHEN, [f"участок г.{(city.name_gorod or '').strip()}", _today()],
                  [Table(cols, [Section(rows=rows)], title=""),
                   Table(scols, [Section(rows=srows)], total=_sum_row(scols, srows, "Итого по городу"))],
                  landscape=True)


# Объекты, которые старая программа не включала в список по котельным
# (potreb_koteln.prg): Рыбница — л/№ 173, 276, 277; Дубоссары — 263–275;
# Каменка — 338–341.
KOTELN_EXCLUDE = {1: {173, 276, 277}, 2: set(range(263, 276)), 3: set(range(338, 342))}


def potreb_koteln(city, opts):
    """«Печать → Потребители по котельным» (potreb_koteln.prg)."""
    names = _koteln()
    skip = KOTELN_EXCLUDE.get(city.cod_gorod, set())
    cols = [Col("Лиц.№", "int"), Col("Наименование потребителя"), Col("Адрес"),
            Col("Тепловая нагр. Гкал/час", "num", 4, total=True), Col("Примечание")]
    sections = []
    for p in Potreb.objects.filter(cod_gor=city.cod_gorod).order_by("id_koteln", "potreb_id"):
        if p.id_koteln not in names or p.potreb_id in skip:
            continue
        if not sections or sections[-1].kot != p.id_koteln:
            s = Section(f"Участок — {names[p.id_koteln]}")
            s.kot = p.id_koteln
            sections.append(s)
        sections[-1].rows.append([p.potreb_id, _name_glav(p), p.adres_sd, n(p.q_otoppr), ""])
    for s in sections:
        s.foot = _sum_row(cols, s.rows, f"всего: {names[s.kot]}")
    scols = [Col("Список котельных"), Col("Тепловая нагр. Гкал/час", "num", 4, total=True)]
    srows = [[names[s.kot], s.foot[3]] for s in sections]
    return Report(FILIAL, [f"участок г.{(city.name_gorod or '').strip()}", _today()],
                  [Table(cols, sections), Table(scols, [Section(rows=srows)], total=_sum_row(scols, srows, "ИТОГО:"))],
                  landscape=True)


KATEGOR_COLS = [(1, "Хозрасчет"), (2, "Местный бюджет"), (3, "Респуб-ий бюджет"),
                (4, "Население бюджет"), (6, "Население абонентское")]


def potreb_kategor(city, opts):
    """«Печать → Нагрузки по категориям» (potreb_kategor.prg): часовая
    нагрузка на отопление по участкам и категориям. Категории 5 и 7 в
    старом отчете не показаны и во «Всего» не входят."""
    sums = {}
    for p in Potreb.objects.exclude(q_otoppr=None).only("cod_gor", "id_kategor", "q_otoppr"):
        key = (p.cod_gor_id, p.id_kategor)
        sums[key] = sums.get(key, Decimal(0)) + p.q_otoppr
    cols = [Col("Участок")] + [Col(t, "num", 5, total=True) for _, t in KATEGOR_COLS] + [Col("Всего по участку", "num", 5, total=True)]
    rows = []
    for c in Danrasch.objects.exclude(cod_gorod=None).order_by("cod_gorod"):
        vals = [sums.get((c.cod_gorod, k), Decimal(0)) for k, _ in KATEGOR_COLS]
        rows.append([f"г.{(c.name_gorod or '').strip()}"] + vals + [sum(vals, Decimal(0))])
    return Report("Часовые нагрузки по категориям потребителей", [_today()],
                  [Table(cols, [Section(rows=rows)], total=_sum_row(cols, rows, "Всего по категориям потреб."))])


# --- Отчеты, добавленные в программе 2026 года --------------------------------

def _tarif(city):
    from .models import TarifGorod
    return {k: (n(z), n(d)) for k, (z, d) in TarifGorod.dlya(city).items()}


POLEZ_COLS = [Col("Категории потребителей"),
              Col("Ведомст. жильё", "num", 2, total=True),
              Col("Население абонентское (с господд.)", "num", 2, total=True),
              Col("Население абонентское (без господд.)", "num", 2, total=True),
              Col("Население хозрасчет", "num", 2, total=True),
              Col("Всего население (Гкал)", "num", 2, total=True),
              Col("Всего хозрасчёт (Гкал)", "num", 2, total=True),
              Col("Местный", "num", 2, total=True),
              Col("Республ.", "num", 2, total=True),
              Col("Организация СКиИОД", "num", 2, total=True),
              Col("Всего бюджет (Гкал)", "num", 2, total=True),
              Col("ВСЕГО (Гкал)", "num", 2, total=True)]
# Колонки «Всего…», которые в строках «П/У» и «Нагр.» старый отчет оставляет пустыми.
POLEZ_VSEGO = (5, 10, 11)


def polez_rows(city, recs, vid=None):
    """Строки П/У и Нагр. полезного отпуска участка city по записям recs
    (уже отобранным действующим). Подпись и Гкал по колонкам POLEZ_COLS."""
    k = _polez_doli(city)
    rows = []
    for label, row_key in (("П/У", "pu"), ("Нагр.", "nagr")):
        row = [label]
        for col in range(1, len(POLEZ_COLS)):
            row.append(sum((_polez_val(p, vid) * f for p in recs
                            if _polez_row(p) == row_key and (f := _polez_part(p, col, k)) is not None),
                           Decimal(0)))
        rows.append(row)
    return rows


def polez_data(vid=None):
    """Полезный отпуск по участкам: [(код участка, название, [строка П/У,
    строка Нагр.])], строка — подпись и Гкал по колонкам POLEZ_COLS. Это же
    кладется в архив месяца при переходе на новый месяц."""
    return [(c.cod_gorod, (c.name_gorod or "").strip(),
             polez_rows(c, list(_active(Potreb.objects.filter(cod_gor=c.cod_gorod))), vid))
            for c in Danrasch.objects.exclude(cod_gorod=None).order_by("cod_gorod")]


def polez_otpusk(city, opts):
    """Кнопка «Полезный отпуск» (polez_otpusk.frx): отпуск тепловой энергии по
    категориям потребителей по всем участкам — по приборам учета (П/У,
    tip_rasot=1) и по нагрузке (tip_rasot 2, 3). Бюджет местный и
    республиканский — по одной колонке вместе с дизельным топливом, как в
    последнем варианте старого отчета; население абонентское — на «с
    господдержкой» и «без» в пропорции Гкал, указанных в исходных данных
    участка. Сложение и отборы повторяют старую программу: хозрасчёт по
    прибору берется без учета вида топлива, а дизельное топливо для
    населения и СКиИОД в отчет не входит. Оформление — как в старом отчете:
    строки П/У и Нагр. по каждому городу, «Всего по г.», в конце П/У, Нагр. и
    «Итого по филиалу»."""
    vid = opts.get("vid")

    def shown(row):
        return [None if i in POLEZ_VSEGO else v for i, v in enumerate(row)]

    def links(gorod, row_key, row):
        return [None if i == 0 or row[i] is None else _polez_url(gorod, row_key, i, vid) for i in range(len(row))]

    sections, pu_rows, nagr_rows = [], [], []
    for gorod, name, rows in polez_data(vid):
        sec = Section(rows=[shown(r) for r in rows])
        sec.links = [links(gorod, key, r) for key, r in zip(("pu", "nagr"), sec.rows)]
        sec.foot = _sum_row(POLEZ_COLS, rows, f"Всего по г.{name}")
        sec.foot_links = links(gorod, "all", sec.foot)
        sections.append(sec)
        pu_rows.append(rows[0])
        nagr_rows.append(rows[1])
    pu, nagr = _sum_row(POLEZ_COLS, pu_rows, "П/У"), _sum_row(POLEZ_COLS, nagr_rows, "Нагр.")
    itog = Section(rows=[shown(pu), shown(nagr)])
    itog.links = [links(0, key, r) for key, r in zip(("pu", "nagr"), itog.rows)]
    itog.foot = _sum_row(POLEZ_COLS, [pu, nagr], "Итого по филиалу")
    itog.mark_foot = True
    itog.foot_links = links(0, "all", itog.foot)
    tar = _tarif(city)  # тарифы выбранного участка

    def t(i, d=False):
        return tar.get(i, (Decimal(0), Decimal(0)))[1 if d else 0]
    tarif_row = ["Тариф (руб)", t(9), t(12), t(11), t(9), t(9), t(1), t(2), t(2), t(10), None, None]
    head = [[("Категории потребителей", 1, 2), ("Население (Гкал)", 4, 1), ("Всего население (Гкал)", 1, 2),
             ("Всего хозрасчёт (Гкал)", 1, 2), ("Бюджет (Гкал)", 2, 1), ("Организация СКиИОД", 1, 2),
             ("Всего бюджет (Гкал)", 1, 2), ("ВСЕГО (Гкал)", 1, 2)],
            [("Ведомст. жильё", 1, 1), ("Население абонентское (с господд.)", 1, 1),
             ("Население абонентское (без господд.)", 1, 1), ("Население хозрасчет", 1, 1),
             ("Местный", 1, 1), ("Республ.", 1, 1)]]
    lines = ["тепловой энергии по категориям потребителей",
             f"за {(city.name_mes or '').strip()} {city.god} г."] + ([dict(VID_USLUGA)[vid]] if vid else [])
    return Report("ПОЛЕЗНЫЙ ОТПУСК", lines,
                  [Table(POLEZ_COLS, [Section(rows=[tarif_row], mark_rows=True)] + sections + [itog], head=head)],
                  signatures=[_ved_spec(city, "\u2003" * 12)],
                  landscape=True,
                  approve=_utverjdau(city))


def _ved_spec(city, sep):
    from .models import Podpisant
    dolj, fio = _signer(Podpisant.VED_SPEC, city)
    return f"{dolj or 'Ведущий специалист'}{sep}{fio}"


def _utverjdau(city):
    """Гриф «УТВЕРЖДАЮ» из справочника подписантов."""
    from .models import Podpisant
    dolj, fio = _signer(Podpisant.UTVERJDAU, city)
    return ["УТВЕРЖДАЮ:"] + ([dolj] if dolj else []) + [f"______________ {fio}", _approve()[1]]


def _polez_val(p, vid):
    """Что суммирует полезный отпуск по записи при выбранной услуге."""
    if vid == "otopl":
        return n(p.nach_otopl)
    if vid == "gvs":
        return _gvs(p)
    return n(p.nach_otopl) + _gvs(p)


def _polez_row(p):
    """Строка отчета: П/У — расчет по прибору, Нагр. — по нагрузке."""
    return "pu" if p.tip_rasot == 1 else "nagr" if p.tip_rasot in (2, 3) else None


def _polez_doli(city):
    """Доли населения абонентского с господдержкой и без нее."""
    gpod, gbez = n(city.gkl_pod), n(city.gkl_bezpod)
    if gpod + gbez > 0:
        return gpod / (gpod + gbez), gbez / (gpod + gbez)
    return Decimal(0), Decimal(0)


def _polez_part(p, col, doli):
    """Множитель, с которым запись входит в колонку col полезного отпуска,
    или None, если не входит. Правила — как в кнопке «Полезный отпуск»
    старой программы (см. polez_otpusk)."""
    kat, diz = p.id_kategor, bool(p.vid_topl)
    rules = {
        1: kat == 4 and not diz,
        2: kat == 6 and not diz,
        3: kat == 6 and not diz,
        4: kat == 5 and not diz,
        5: kat in (4, 5, 6) and not diz,
        6: kat == 1 and (p.tip_rasot == 1 or not diz),
        7: kat == 2,
        8: kat == 3,
        9: kat == 7 and not diz,
        10: kat in (2, 3),
        11: not diz or kat in (2, 3),
    }
    if not rules.get(col):
        return None
    return {2: doli[0], 3: doli[1]}.get(col, Decimal(1))


POLEZ_ROWS = {"pu": "П/У", "nagr": "Нагр.", "all": "П/У и Нагр."}


def _polez_url(gorod, row_key, col, vid):
    from urllib.parse import urlencode
    from django.urls import reverse
    return reverse("abonents:polez_rasshifr") + "?" + urlencode(
        {"uch": gorod, "row": row_key, "col": col, "vid": vid or ""})


def polez_rasshifr(params):
    """Расшифровка ячейки полезного отпуска: записи, из которых она сложилась."""
    try:
        col = int(params.get("col", ""))
        uch = int(params.get("uch", "0"))
    except ValueError:
        return None
    row_key = params.get("row")
    vid = params.get("vid") or ""
    if not 1 <= col < len(POLEZ_COLS) or row_key not in POLEZ_ROWS or vid not in dict(VID_USLUGA):
        return None
    cities = {c.cod_gorod: c for c in Danrasch.objects.exclude(cod_gorod=None)}
    if uch and uch not in cities:
        return None
    qs = _active(Potreb.objects.filter(cod_gor__in=[uch] if uch else list(cities)))
    names, sposob = _kategor(), dict(TIP_RASOT)
    cols = [Col("Участок"), Col("Л/№", "int"), Col("№ дог.", "int"), Col("Наименование"), Col("Объект"),
            Col("Категория"), Col("Способ расчета"), Col("Диз. топливо"),
            Col("Отопление", "num", 3, total=True), Col("ГВС", "num", 3, total=True),
            Col("Потери ГВС", "num", 3, total=True), Col("Доля", "num", 4), Col("В ячейку", "num", 3, total=True)]
    rows, links = [], []
    for p in qs.order_by("cod_gor", "id_kategor", "potreb_id"):
        if row_key != "all" and _polez_row(p) != row_key or _polez_row(p) is None:
            continue
        f = _polez_part(p, col, _polez_doli(cities[p.cod_gor_id]))
        v = _polez_val(p, vid)
        if f is None or v == 0:
            continue
        rows.append([(cities[p.cod_gor_id].name_gorod or "").strip(), p.potreb_id, p.nom_dog, p.name_potre,
                     p.glav_potre, names.get(p.id_kategor, p.id_kategor), sposob.get(p.tip_rasot, p.tip_rasot),
                     "да" if p.vid_topl else "", n(p.nach_otopl) if vid != "gvs" else Decimal(0),
                     n(p.nach_gvs) if vid != "otopl" else Decimal(0), n(p.poter_gvs) if vid != "otopl" else Decimal(0),
                     f if col in (2, 3) else None, v * f])
    title = {7: "Местный бюджет", 8: "Республиканский бюджет"}.get(col, POLEZ_COLS[col].title)
    lines = [f"Колонка: {title}", f"Строка: {POLEZ_ROWS[row_key]}",
             "Участок: " + ((cities[uch].name_gorod or "").strip() if uch else "все (итого по филиалу)")]
    if vid:
        lines.append(dict(VID_USLUGA)[vid])
    total = _sum_row(cols, rows, "Итого")
    total[-2] = None
    return Report("Расшифровка: полезный отпуск", lines, [Table(cols, [Section(rows=rows)], total=total)],
                  notes=[f"Записей: {len(rows)}"], landscape=True)


NAG_MINIST = (1, 2, 3, 4, 7, 8, 9, 10, 19)


def nag_minist(city, opts):
    """Кнопка «Нагрузка по министерствам» (nag_minist1): часовая нагрузка
    республиканского бюджета по министерствам. Как в старой программе,
    отбор идет по всем участкам и только по министерствам из списка."""
    qs = _active(Potreb.objects.filter(nom_dog__gt=0, id_kategor=3, cod_minist__in=NAG_MINIST)
                 .exclude(id_katabon=3)).order_by("cod_minist", "potreb_id")
    names = dict(Minist.objects.values_list("cod_minist", "name_minis"))
    cols = [Col("№ п/п", "int"), Col("Абонент"), Col("ГВС", "num", 5, total=True),
            Col("Отопление", "num", 5, total=True), Col("Всего", "num", 5, total=True)]
    sections = []
    for s in _grouped(qs, lambda p: p.cod_minist_id, lambda k: (names.get(k) or "").strip(),
                      lambda p: [0, p.name_potre, n(p.q_gvs), n(p.q_otoppr), n(p.q_gvs) + n(p.q_otoppr)], cols):
        s.foot = _sum_row(cols, s.rows, "ВСЕГО:", label_col=1)
        sections.append(s)
    rows = [r for s in sections for r in s.rows]
    return Report("РАСПРЕДЕЛЕННАЯ НАГРУЗКА ПО МИНИСТЕРСТВАМ", [f"{(city.name_mes or '').strip()} {city.god} г."],
                  [Table(cols, sections, total=_sum_row(cols, rows, "ИТОГО:", label_col=1))],
                  signatures=[_ved_spec(city, f" {LINE} ")])


def _ispol_sign(city):
    from .models import Podpisant
    dolj, fio = _signer(Podpisant.PRIBORY, city)
    return [f"{dolj}", f"Подпись _______________ дата ___________      {fio}"]


def _pribor_head(city, title):
    return ([title, f"по участку г.{(city.name_gorod or '').strip()}",
             f"за {(city.name_mes or '').strip()} {city.god} г."],
            _approve() + ["Приложение №2", "к приказу от 08.04.2024г.№290"])


def otch_pribor(city, opts):
    """Кнопка «Отчет по приборам учета (отопление)» (otch_potr_pribor):
    потребление по разнице показаний с коэффициентом прибора учета (если
    прибор не указан — коэффициент 1). Абонент с несколькими приборами
    показан строкой на каждый прибор."""
    devlist = {d.kod: d for d in SetUrdevicelist.objects.exclude(kod=None)}
    devices = {}
    for d in SetUrdevice.objects.exclude(p_kod=None):
        try:
            devices.setdefault(round(float(d.p_kod)), []).append(d)
        except ValueError:
            continue
    cols = [Col("Л/№", "int"), Col("Наименование абонента (потребителя)"), Col("Фактический адрес объекта"),
            Col("Тип и номер прибора"), Col("ед. изм."), Col("Показания предыдущие", "num"),
            Col("Показания настоящие", "num"), Col("Расчет потребления (Гкал)", "num", 3, total=True),
            Col("Примечание")]
    rows = []
    for p in _active(Potreb.objects.filter(cod_gor=city.cod_gorod)).order_by("potreb_id"):
        for d in devices.get(p.potreb_id) or [None]:
            info = devlist.get(d.d_kod) if d else None
            koef = info.koef if info and info.koef is not None else Decimal(1)
            rasch = (n(p.pok_new) - n(p.pok_old)) * koef
            if rasch <= 0:
                continue
            tip = f"{(info.name or '').strip() if info else ''}  №{(d.d_num or '').strip() if d else ''}"
            rows.append([p.potreb_id, p.glav_potre, p.adres_sd, tip, (info.edn or "").strip() if info else "",
                         n(p.pok_old), n(p.pok_new), rasch, p.prim_chot])
    lines, approve = _pribor_head(city, "Отчет о потреблении тепловой энергии по приборам учета тепловой энергии")
    return Report(lines[0], lines[1:], [Table(cols, [Section(rows=rows)], total=_sum_row(cols, rows, "Итого"))],
                  signatures=_ispol_sign(city), landscape=True, approve=approve)


def otch_priborg(city, opts):
    """Кнопка «Отчет по приборам учета (ГВС)» (otch_potr_priborg): потребление
    по разнице показаний счетчика ГВС, 1 куб.м = 0,0533 Гкал."""
    from .raschet import GVS_GKAL
    cols = [Col("Л/№", "int"), Col("Наименование абонента (потребителя)"), Col("Фактический адрес объекта"),
            Col("ед. изм."), Col("Показания предыдущие", "num"), Col("Показания настоящие", "num"),
            Col("Расчет потребления (Гкал)", "num", 3, total=True), Col("Примечание")]
    rows = []
    for p in _active(Potreb.objects.filter(cod_gor=city.cod_gorod)).order_by("potreb_id"):
        rasch = (n(p.pok_newg) - n(p.pok_oldg)) * GVS_GKAL
        if rasch > 0:
            rows.append([p.potreb_id, p.glav_potre, p.adres_sd, "куб.м.", n(p.pok_oldg), n(p.pok_newg), rasch,
                         p.prim_chgvs])
    lines, approve = _pribor_head(city, "Отчет о потреблении ГВС по приборам учета")
    return Report(lines[0], lines[1:], [Table(cols, [Section(rows=rows)], total=_sum_row(cols, rows, "Итого"))],
                  signatures=_ispol_sign(city), landscape=True, approve=approve)


DATE = "date"


def device_pov(city, opts):
    """Кнопка «Поверка приборов» (device_pov): приборы учета, у которых
    дата следующей поверки раньше выбранной даты."""
    try:
        date = datetime.date.fromisoformat(opts.get("date") or "")
    except ValueError:
        date = datetime.date.today()
    qs = Potreb.objects.filter(cod_gor=city.cod_gorod)
    kat = opts.get("kateg")
    if kat:
        qs = qs.filter(id_kategor=int(kat))
    names = dict(qs.values_list("potreb_id", "name_potre"))
    rows = []
    for d in SetUrdevice.objects.filter(d_next_pov__lt=date).exclude(p_kod=None).order_by("d_pov", "pk"):
        try:
            pid = int(float(d.p_kod))
        except ValueError:
            continue
        if pid in names:
            rows.append([pid, names[pid], d.d_pov, d.d_next_pov])
    cols = [Col("№ Лиц.сч.", "int"), Col("Наименование потребителя"), Col("Дата поверки", "date"),
            Col("Дата следующей поверки", "date")]
    lines = [f"г. {(city.name_gorod or '').strip()}"]
    if kat:
        lines.append(_kategor().get(int(kat)) or "")
    lines.append(f"следующая поверка ранее {date:%d.%m.%Y}")
    return Report("Список приборов с истекающим (истекшим) сроком поверки", lines,
                  [Table(cols, [Section(rows=rows)])], notes=[f"Всего: {len(rows)}"])


# --- Отчеты за период по закрытым месяцам (форма arhiv) ----------------------

MONTH = "month"     # параметр «месяц»: значение ГГГГ-ММ
NUMBER = "number"   # параметр «число» (лицевой №)


def _ym(value):
    g, m = value.split("-")
    return int(m), int(g)


def _period(opts):
    from . import arhiv
    nach, kon = _ym(opts["s"]), _ym(opts["po"])
    if (nach[1], nach[0]) > (kon[1], kon[0]):
        nach, kon = kon, nach
    spisok = arhiv.mesyacy(nach, kon)
    text = arhiv.nazvanie(*nach) + ("" if nach == kon else " — " + arhiv.nazvanie(*kon))
    return spisok, text


def _net_arhiva(missing):
    return [f"Нет архива за: {', '.join(missing)}"] if missing else []


def arh_polez(city, opts):
    """Полезный отпуск за период (polez_nak): по каждому закрытому месяцу и
    участку — те же колонки, что в «Полезном отпуске» (П/У и нагрузка вместе),
    итог месяца и итог за период."""
    from . import arhiv
    vid = opts.get("vid") or None
    spisok, text = _period(opts)
    cols = [Col("Месяц, участок")] + POLEZ_COLS[1:]
    sections, itogi, missing = [], [], []
    for mes, god, m in arhiv.zagruzit(spisok):
        if m is None:
            missing.append(arhiv.nazvanie(mes, god))
            continue
        rows = []
        for d in sorted(m.danrasch, key=lambda d: d.cod_gorod or 0):
            if d.cod_gorod is None:
                continue
            pu, nagr = polez_rows(d, m.dejstvuyushchie(d.cod_gorod), vid)
            rows.append([(d.name_gorod or "").strip()] + [a + b for a, b in zip(pu[1:], nagr[1:])])
        sec = Section(title=m.nazvanie, rows=rows)
        sec.foot = _sum_row(cols, rows, f"Итого за {m.nazvanie}")
        itogi.append(sec.foot)
        sections.append(sec)
    total = _sum_row(cols, itogi, "Итого за период")
    lines = ["тепловой энергии по категориям потребителей", f"за {text}"] + \
        ([dict(VID_USLUGA)[vid]] if vid else [])
    return Report("ПОЛЕЗНЫЙ ОТПУСК ЗА ПЕРИОД", lines, [Table(cols, sections, total=total)],
                  notes=_net_arhiva(missing), signatures=[_ved_spec(city, "\u2003" * 12)],
                  landscape=True, approve=_utverjdau(city))


def arh_ras_ab(city, opts):
    """Расчет количества тепловой энергии, отпущенной абоненту, за период
    (ras_abonenta, формы arhiv и fak_otp): по месяцам из архива."""
    from . import arhiv
    spisok, text = _period(opts)
    try:
        ls = int(opts.get("ls") or 0)
    except ValueError:
        ls = 0
    cols = [Col("Период потребления"), Col("Фактич. темп-ра", "num", 1), Col("Темпер. коэф.", "num", 3),
            Col("Теплов. нагр.", "num", 4), Col("Перерасч. Гкал", "num", 3, total=True),
            Col("Кол. дней", "num", 0), Col("Начисл. нагруз.", "num", 3, total=True),
            Col("Потери нагруз.", "num", 3, total=True), Col("Начисл. прибор", "num", 4, total=True),
            Col("Потери прибор", "num", 3, total=True), Col("Отопл. Гкал", "num", 3, total=True),
            Col("ГВС Гкал", "num", 3, total=True), Col("Всего Гкал", "num", 3, total=True)]
    rows, missing, head = [], [], None
    for mes, god, m in arhiv.zagruzit(spisok):
        p = m.licevoj(ls) if m else None
        if p is None:
            missing.append(arhiv.nazvanie(mes, god))
            continue
        head = p
        d = m.gorod(p.cod_gor_id)
        rows.append([m.nazvanie, d.t_vozd if d else None, d.tkoef if d else None, p.q_otoppr, n(p.gkl_perer),
                     p.dnotop, n(p.nach_otnag), n(p.poter_nag), n(p.nach_otpr), n(p.poter_pr),
                     n(p.nach_otopl), n(p.nach_gvs), n(p.nach_otopl) + n(p.nach_gvs)])
    head = Potreb.objects.filter(potreb_id=ls).first() or head
    lines = []
    if head is not None:
        lines.append(f"{(head.name_potre or '').strip()} ({(head.glav_potre or '').strip()})")
        lines.append(f"по договору № {head.nom_dog or ''} от {head.data_dog:%d.%m.%Y}" if head.data_dog
                     else f"по договору № {head.nom_dog or ''}")
    lines.append(f"Л/№ {ls}, {text}" if ls else "Укажите лицевой №")
    notes = [f"Нет лицевого № {ls} в архиве за: {', '.join(missing)}"] if ls and missing else []
    return Report("РАСЧЕТ количества тепловой энергии, отпущенной абоненту", lines,
                  [Table(cols, [Section(rows=rows)], total=_sum_row(cols, rows, "ИТОГО:"))],
                  notes=notes, signatures=[_ved_spec(city, "\u2003" * 12)], landscape=True)


BUDJ_HEAD = [[("Месяц, участок", 1, 2), ("Вед. жильё (хозрасчёт)", 2, 1), ("Вед. жильё (бюджет)", 2, 1),
              ("Абонентское население", 2, 1), ("Хозрасчёт", 1, 2), ("Местный б.", 1, 2),
              ("Республ. б.", 1, 2), ("Всего", 1, 2)],
             [("Отопление", 1, 1), ("ГВС", 1, 1)] * 3]


def budj_row(d, recs):
    """Строка ведомости в разрезе бюджетов по участку d (budj_nak.frx):
    отопление и ГВС ведомственного жилья и абонентского населения отдельно,
    хозрасчёт и бюджеты — всего. Суммы — те же, что в «Полезном отпуске»
    (П/У и нагрузка вместе)."""
    def vse(vid):
        pu, nagr = polez_rows(d, recs, vid)
        return [a + b for a, b in zip(pu[1:], nagr[1:])]
    o, g, v = vse("otopl"), vse("gvs"), vse(None)
    # индексы POLEZ_COLS без первой колонки: 0 ведомств. жильё (кат. 4), 3 нас. хозрасчет (5),
    # 4 всего население (4+5+6), 5 хозрасчёт, 6 местный, 7 республ., 10 ВСЕГО
    nas = [x[4] - x[0] - x[3] for x in (o, g)]
    return [o[3], g[3], o[0], g[0], nas[0], nas[1], v[5], v[6], v[7], v[10]]


def arh_budj(city, opts):
    """Ведомость фактического отпуска т/энергии в разрезе бюджетов за период
    (budj_nak, вариант «по категориям с ГВС» формы arhiv): по каждому
    закрытому месяцу и участку, итог месяца и итог за период. В старой
    программе ведомость брала готовые итоги из Arhiv\\OtММГГГГ.dbf, здесь они
    считаются по сохранённым записям месяца, как «Полезный отпуск за период»."""
    from . import arhiv
    spisok, text = _period(opts)
    cols = [Col("Месяц, участок")] + [Col("", "num", 2, total=True) for _ in range(10)]
    sections, itogi, missing = [], [], []
    for mes, god, m in arhiv.zagruzit(spisok):
        if m is None:
            missing.append(arhiv.nazvanie(mes, god))
            continue
        rows = [[(d.name_gorod or "").strip()] + budj_row(d, m.dejstvuyushchie(d.cod_gorod))
                for d in sorted(m.danrasch, key=lambda d: d.cod_gorod or 0) if d.cod_gorod is not None]
        sec = Section(title=m.nazvanie, rows=rows)
        sec.foot = _sum_row(cols, rows, "Итого за месяц")
        itogi.append(sec.foot)
        sections.append(sec)
    total = _sum_row(cols, itogi, "Итого за период по филиалу")
    return Report("ВЕДОМОСТЬ ФАКТИЧЕСКОГО ОТПУСКА Т/ЭНЕРГИИ В РАЗРЕЗЕ БЮДЖЕТОВ",
                  ["по МГУП «Тирастеплоэнерго» (Гкал)", f"за {text}"],
                  [Table(cols, sections, total=total, head=BUDJ_HEAD)],
                  notes=_net_arhiva(missing), signatures=[_ved_spec(city, "\u2003" * 12)],
                  landscape=True, approve=_utverjdau(city))


DOG_MAX_MES = 12


def _obekt(p):
    adr = " ".join(x for x in ((p.adres_sd or "").strip(), (p.oba_dom or "").strip(), (p.oba_kv or "").strip()) if x)
    name = (p.nameobakt or "").strip() or (p.name_potre or "").strip()
    return f"{name}, {adr}" if adr else name


def arh_dog(city, opts):
    """Расшифровка отпуска по договору за период (ved_budjet.frx, форма
    arhiv): объекты договора по категориям, Гкал (отопление + ГВС) по
    месяцам и всего. Без номера договора — все лицевые участка с отпуском:
    так же старая программа выгружала «Начисление юрлица» в Excel."""
    from . import arhiv
    spisok, text = _period(opts)
    notes = []
    if len(spisok) > DOG_MAX_MES:
        notes.append(f"Показаны первые {DOG_MAX_MES} месяцев периода.")
        spisok = spisok[:DOG_MAX_MES]
    try:
        dog = int(opts.get("dog") or 0)
    except ValueError:
        dog = 0
    cols = [Col("Наименование объекта"), Col("Л/№", "int"), Col("№ дог.", "int")] + \
        [Col(f"{m:02d}.{g}", "num", 3, total=True) for m, g in spisok] + [Col("ВСЕГО", "num", 3, total=True)]
    data, missing = {}, []
    for i, (mes, god, m) in enumerate(arhiv.zagruzit(spisok)):
        if m is None:
            missing.append(arhiv.nazvanie(mes, god))
            continue
        for p in m.potreb:
            if p.cod_gor_id != city.cod_gorod or (dog and p.nom_dog != dog):
                continue
            row = data.setdefault(p.potreb_id, {"p": p, "nac": [Decimal(0)] * len(spisok)})
            row["p"] = p
            row["nac"][i] = n(p.nach_otopl) + n(p.nach_gvs)
    items = [r for r in data.values() if dog or any(r["nac"])]
    items.sort(key=lambda r: (r["p"].id_kategor or 0, r["p"].nom_dog or 0, r["p"].potreb_id or 0))
    kateg = _kategor()
    sections, key = [], object()
    for r in items:
        p = r["p"]
        if p.id_kategor != key:
            key = p.id_kategor
            sections.append(Section(title=(kateg.get(p.id_kategor) or "").strip()))
        sections[-1].rows.append([_obekt(p), p.potreb_id, p.nom_dog] + r["nac"] + [sum(r["nac"])])
    for sec in sections:
        sec.foot = _sum_row(cols, sec.rows, f"Итого: {sec.title}" if sec.title else "Итого")
    total = _sum_row(cols, [x for sec in sections for x in sec.rows], "ВСЕГО")
    lines = _city_lines(city, month=False)
    if dog:
        glav = next((r["p"] for r in items), None) or Potreb.objects.filter(cod_gor=city.cod_gorod, nom_dog=dog).first()
        if glav is not None:
            lines.append((glav.glav_potre or glav.name_potre or "").strip())
    lines.append(f"за {text}")
    if dog and not items:
        notes.append(f"Договора № {dog} нет в архиве за выбранный период.")
    title = f"РАСШИФРОВКА ОТПУСКА ПО ДОГОВОРУ № {dog}" if dog else "РАСШИФРОВКА ОТПУСКА ПО ЛИЦЕВЫМ"
    return Report(title, lines, [Table(cols, sections, total=total)],
                  notes=notes + _net_arhiva(missing), signatures=[_ved_spec(city, "\u2003" * 12)], landscape=True)


def arh_potreb_god(city, opts):
    """Данные о фактическом потреблении тепловой энергии за год (potreb_god):
    бюджетные потребители (местный и республиканский бюджет, кроме категории
    абонента 3) участка, начислено всего по месяцам. В старой программе год
    в коде был зашит (2009), здесь — выбранный."""
    from catalogs.models import Minist
    from . import arhiv
    try:
        god = int(opts.get("god") or 0)
    except ValueError:
        god = 0
    cols = [Col("Наименование потребителя"), Col("№ догов.", "int"), Col("Дата заключ.", "date")] + \
        [Col(m[:3] + ".", "num", 3, total=True) for m in arhiv.MES] + [Col("Итого за год", "num", 3, total=True)]
    data, missing = {}, []
    for mes, _, m in arhiv.zagruzit([(i, god) for i in range(1, 13)]):
        if m is None:
            missing.append(arhiv.MES[mes - 1])
            continue
        for p in m.potreb:
            if p.cod_gor_id != city.cod_gorod or p.id_kategor not in (2, 3) or p.id_katabon == 3:
                continue
            row = data.setdefault(p.potreb_id, {"p": p, "nac": [Decimal(0)] * 12})
            row["p"] = p
            row["nac"][mes - 1] = n(p.nac_vsego)
    kateg = _kategor()
    minist = dict(Minist.objects.values_list("cod_minist", "name_minis"))
    items = sorted(data.values(), key=lambda r: (r["p"].id_kategor or 0, r["p"].cod_minist_id or 0, r["p"].potreb_id))
    sections, current, key = [], None, None
    for r in items:
        p = r["p"]
        k = (p.id_kategor, p.cod_minist_id)
        if k != key:
            title = (kateg.get(p.id_kategor) or "").strip()
            if p.id_kategor == 3:
                title += f" — {(minist.get(p.cod_minist_id) or '').strip()}"
            current = Section(title=title)
            sections.append(current)
            key = k
        current.rows.append([(p.name_potre or "").strip(), p.nom_dog, p.data_dog] + r["nac"] + [sum(r["nac"])])
    for sec in sections:
        sec.foot = _sum_row(cols, sec.rows, f"Итого: {sec.title}")
    total = _sum_row(cols, [r for sec in sections for r in sec.rows], "ИТОГО:")
    return Report(f"ДАННЫЕ О ФАКТИЧЕСКОМ ПОТРЕБЛЕНИИ ТЕПЛОВОЙ ЭНЕРГИИ ЗА {god} г.", _city_lines(city, month=False),
                  [Table(cols, sections, total=total)],
                  notes=[f"Нет архива за: {', '.join(missing)} {god}"] if missing else [],
                  signatures=[_ved_spec(city, "\u2003" * 12)], landscape=True)


def _arh_mes_default():
    from . import arhiv
    last = arhiv.poslednij()
    if last is None:
        today = datetime.date.today()
        return f"{today.year:04d}-{today.month:02d}"
    return f"{last[1]:04d}-{last[0]:02d}"


def _arh_gody():
    from . import arhiv
    return [(str(g), str(g)) for g in arhiv.gody()] or [(str(datetime.date.today().year), str(datetime.date.today().year))]


# --- Ведомости, реестры -------------------------------------------------------

def reestr_dog(city, opts):
    """«Ведомости, реестры → Реестр действующих договоров» (proced.reestr_dog,
    reestr_dog.frx): договоры участка с заключённым или возобновлённым
    статусом по категориям потребителей — абонент, № и дата договора. На
    договор одна строка (главный лицевой: абонент, затем меньший л/№)."""
    names = _kategor()
    recs = _active(Potreb.objects.filter(cod_gor=city.cod_gorod, nom_dog__gt=0)).order_by("nom_dog", "potreb_id")
    dog = {}
    for p in recs:
        cur = dog.get(p.nom_dog)
        if cur is None or (cur.id_katabon != 1 and p.id_katabon == 1):
            dog[p.nom_dog] = p
    items = sorted(dog.values(), key=lambda p: (p.id_kategor or 0, p.nom_dog))
    cols = [Col("№ п/п", "int"), Col("Абонент"), Col("Номер договора", "int"), Col("Дата заключения договора", "date")]
    sections, cur, kat = [], None, object()
    for p in items:
        if p.id_kategor != kat:
            kat = p.id_kategor
            cur = Section(f"Категория потребителя: {names.get(kat) or kat or '—'}")
            sections.append(cur)
        cur.rows.append([len(cur.rows) + 1, (p.name_potre or "").strip(), p.nom_dog, p.data_dog])
    for s in sections:
        s.foot = [None, f"договоров: {len(s.rows)}", None, None]
    return Report("Реестр действующих договоров на отпуск тепловой энергии",
                  [f"г. {(city.name_gorod or '').strip()}",
                   f"по состоянию на {(city.name_mes or '').strip()} {city.god} г."],
                  [Table(cols, sections, total=[None, f"Всего договоров: {len(items)}", None, None])],
                  signatures=_signatures(city))


def _pisma_mesyacy():
    from .models import Pismo
    mes = sorted(set(Pismo.objects.values_list("god", "nom_mes")), reverse=True)
    from .nov_mes import MONTHS
    return [(f"{g}-{m:02d}", f"{MONTHS[m - 1]} {g}") for g, m in mes] or [("", "писем ещё не отправляли")]


VID_KANAL = [("", "Почта и КРЭДО"), ("mail", "Почта"), ("kredo", "КРЭДО")]


def vedom_pisem(city, opts):
    """«Ведомости, реестры → Ведомость отправленных писем» (proced.vedom,
    vedommail.frx): документы, успешно отправленные абонентам участка за
    месяц, по реестру писем — л/№, № договора, потребитель, дата отправки."""
    from .models import Pismo
    from .nov_mes import MONTHS
    qs = Pismo.objects.filter(cod_gor=city.cod_gorod, status=Pismo.OK)
    mes = opts.get("mes") or ""
    if mes:
        m, g = _ym(mes)
    else:
        m, g = city.nom_mes, city.god
    qs = qs.filter(nom_mes=m, god=g)
    if opts.get("kanal"):
        qs = qs.filter(kanal=opts["kanal"])
    cols = [Col("№ ПП", "int"), Col("Л.№", "int"), Col("№ дог", "int"), Col("Наименование потребителя"),
            Col("Документ"), Col("Куда"), Col("Дата отправки письма")]
    vid = {"schet": "Информация-счёт", "akt": "Акт"}
    rows = [[i, z.potreb_id, z.nom_dog, z.name_potre, vid.get(z.vid, z.vid), z.get_kanal_display(),
             timezone.localtime(z.created).strftime("%d.%m.%Y %H:%M")]
            for i, z in enumerate(qs.order_by("created", "nomer"), 1)]
    mes_txt = f"{MONTHS[m - 1]} {g} г." if m and g else ""
    return Report("ВЕДОМОСТЬ", ["отправленных писем абонентам", f"г. {(city.name_gorod or '').strip()}",
                                f"документы за {mes_txt}" if mes_txt else "",
                                f"сформирована {timezone.localtime().strftime('%d.%m.%Y %H:%M')}"],
                  [Table(cols, [Section(rows=rows)], total=[None, None, None, f"Всего писем: {len(rows)}", None, None, None])])


NAGR_KAT = [(1, "Хозрасчет"), (2, "Местный бюджет"), (3, "Республиканский бюджет"), (4, "Население бюджет"),
            (5, "Население хозрасчет"), (6, "Население абонентское"), (7, "Организации СКиОД")]


def nagr_kot_kat(city, opts):
    """«Ведомости, реестры → Нагрузка котельных по категориям»
    (nagr_kot_kat.prg, в старой программе — книга Excel по шаблону): по
    каждой котельной участка часовая нагрузка действующих договоров на
    отопление (q_otoppr) и ГВС (q_gvs) по категориям потребителей."""
    sums = {}
    for p in _active(Potreb.objects.filter(cod_gor=city.cod_gorod)).only("id_koteln", "id_kategor", "q_otoppr", "q_gvs"):
        k = (p.id_koteln, p.id_kategor)
        o, g = sums.get(k, (Decimal(0), Decimal(0)))
        sums[k] = (o + n(p.q_otoppr), g + n(p.q_gvs))
    kotly = list(GvsKot.objects.filter(kod_gor=city.cod_gorod).order_by("id_kot"))
    known = {k.id_kot for k in kotly}
    # Шапка как в шаблоне «Нагрузка котельных по категориям.xltx»: по каждой
    # категории «отопл.» и «ГВС», без колонок «Всего».
    cols = [Col("Наименование котельной")]
    for _, t in NAGR_KAT:
        cols += [Col(f"{t}: отопл.", "num", 5, total=True), Col(f"{t}: ГВС", "num", 5, total=True)]
    rows = []

    def row(name, kot):
        r = [name]
        for kat, _ in NAGR_KAT:
            r += list(sums.get((kot, kat), (Decimal(0), Decimal(0))))
        return r
    for k in kotly:
        r = row((k.name_kot or "").strip() or f"Котельная {k.id_kot}", k.id_kot)
        if any(r[1:]):
            rows.append(r)
    for kot in sorted({k for k, _ in sums if k not in known}, key=lambda v: (v is None, v or 0)):
        rows.append(row("Котельная не указана" if kot is None else f"Котельная {kot}", kot))
    head = [[("Наименование котельной", 1, 2)] + [(t, 2, 1) for _, t in NAGR_KAT],
            [("отопл.", 1, 1), ("ГВС", 1, 1)] * len(NAGR_KAT)]
    return Report("Нагрузка котельных по категориям потребителей, Гкал/час",
                  [f"г. {(city.name_gorod or '').strip()} данные на {_today()}"],
                  [Table(cols, [Section(rows=rows)], total=_sum_row(cols, rows, "Итого по участку"), head=head)],
                  landscape=True)


VID_PRIBOR = [("nach", "Начислено"), ("nenach", "Не начислено")]
VID_NUL = [("otopl", "Отопление"), ("gvs", "ГВС")]

REPORTS = [
    ReportDef("ot_kateg", "Распределение по бюджетам", "Начисление", ot_kateg,
              options=[("kateg", "Категория", None)], hint="Отпуск тепловой энергии по категориям потребителей"),
    ReportDef("ot_minist", "Расшифровка по министерствам", "Начисление", ot_minist,
              hint="Республиканский бюджет по министерствам"),
    ReportDef("ved_potreb", "Ведомость по приборам", "Начисление", ved_potreb,
              options=[("vid", "Показать", VID_PRIBOR)], hint="Начисление по разнице показаний приборов"),
    ReportDef("nach_nul", "Без начисления", "Начисление", nach_nul,
              options=[("vid", "Услуга", VID_NUL)], hint="Список потребителей без начисления"),
    ReportDef("limit", "Превышение лимита", "Начисление", limit, per_city=False,
              hint="Заканчивающийся лимит (по последнему расчету)"),
    ReportDef("polez_otpusk", "Полезный отпуск", "Начисление", polez_otpusk, per_city=False,
              options=[("vid", "Услуга", VID_USLUGA)], hint="Отпуск тепловой энергии по категориям, все участки"),
    ReportDef("p_kot1", "Потребление по котельным", "Котельные", p_kot1,
              options=[("vid", "Услуга", VID_USLUGA)], hint="Фактическое потребление ГВС и отопления по котельным"),
    ReportDef("p_kot2", "Нагрузка по котельным", "Котельные", p_kot2,
              hint="Распределенная часовая нагрузка по котельным"),
    ReportDef("nag_kateg", "Нагрузка по категориям", "Котельные", nag_kateg,
              hint="Тепловая нагрузка абонентов по категориям"),
    ReportDef("nag_minist", "Нагрузка по министерствам", "Котельные", nag_minist, per_city=False,
              hint="Нагрузка республиканского бюджета по министерствам"),
    ReportDef("otch_pribor", "Потребление по приборам (отопление)", "Приборы учета", otch_pribor,
              hint="Отчет о потреблении по приборам учета тепловой энергии"),
    ReportDef("otch_priborg", "Потребление по приборам (ГВС)", "Приборы учета", otch_priborg,
              hint="Отчет о потреблении ГВС по счетчикам"),
    ReportDef("device_pov", "Поверка приборов", "Приборы учета", device_pov,
              options=[("date", "Поверка ранее", DATE), ("kateg", "Категория", None)],
              hint="Приборы с истекающим (истекшим) сроком поверки"),
    ReportDef("pasp_got", "Паспорта готовности", "Печать", pasp_got,
              hint="Перечень потребителей с паспортами готовности"),
    ReportDef("potreb_koteln", "Потребители по котельным", "Печать", potreb_koteln,
              hint="Перечень потребителей и нагрузка по котельным"),
    ReportDef("potreb_kategor", "Часовые нагрузки по категориям", "Печать", potreb_kategor, per_city=False,
              hint="Нагрузка по участкам и категориям потребителей"),
    ReportDef("reestr_dog", "Реестр действующих договоров", "Ведомости, реестры", reestr_dog,
              hint="Действующие договоры участка по категориям потребителей"),
    ReportDef("vedom_pisem", "Ведомость отправленных писем", "Ведомости, реестры", vedom_pisem,
              options=[("mes", "Месяц", _pisma_mesyacy), ("kanal", "Куда", VID_KANAL)],
              hint="Документы, отправленные абонентам за месяц, по реестру писем"),
    ReportDef("nagr_kot_kat", "Нагрузка котельных по категориям", "Ведомости, реестры", nagr_kot_kat,
              hint="Часовая нагрузка отопления и ГВС по котельным и категориям"),
    ReportDef("arh_polez", "Полезный отпуск за период", "Архив (за период)", arh_polez, per_city=False,
              options=[("s", "С", MONTH), ("po", "по", MONTH), ("vid", "Услуга", VID_USLUGA)],
              hint="По закрытым месяцам, по участкам, с итогом за период"),
    ReportDef("arh_budj", "Отпуск в разрезе бюджетов за период", "Архив (за период)", arh_budj, per_city=False,
              options=[("s", "С", MONTH), ("po", "по", MONTH)],
              hint="Ведомость фактического отпуска: ведомственное жильё и население — отопление и ГВС"),
    ReportDef("arh_dog", "Расшифровка отпуска по договору", "Архив (за период)", arh_dog,
              options=[("dog", "Договор №", NUMBER), ("s", "С", MONTH), ("po", "по", MONTH)],
              hint="Гкал по объектам договора по месяцам; без номера — все лицевые участка (для Excel)"),
    ReportDef("arh_ras_ab", "Расшифровка отпуска по лицевому", "Архив (за период)", arh_ras_ab, per_city=False,
              options=[("ls", "Л/№", NUMBER), ("s", "С", MONTH), ("po", "по", MONTH)],
              hint="Расчет количества тепла, отпущенного абоненту, по месяцам"),
    ReportDef("arh_potreb_god", "Потребление бюджета за год", "Архив (за период)", arh_potreb_god,
              options=[("god", "Год", _arh_gody)],
              hint="Бюджетные потребители: начислено по месяцам года"),
]
BY_KEY = {r.key: r for r in REPORTS}


def options_for(rdef):
    """Варианты выбора для параметров отчета (категории берутся из справочника)."""
    out = []
    for name, label, choices in rdef.options:
        if choices is None:
            choices = [("", "Все категории")] + [(str(k), v) for k, v in sorted(_kategor().items())]
        elif callable(choices):
            choices = choices()
        out.append((name, label, choices))
    return out


def build(rdef, city, params):
    opts = {}
    for name, _, choices in options_for(rdef):
        if choices == DATE:
            value = params.get(name, "")
            try:
                opts[name] = datetime.date.fromisoformat(value).isoformat()
            except ValueError:
                opts[name] = datetime.date.today().isoformat()
            continue
        if choices == MONTH:
            value = params.get(name, "")
            try:
                m, g = _ym(value)
                opts[name] = f"{g:04d}-{m:02d}" if 1 <= m <= 12 and g > 1900 else _arh_mes_default()
            except ValueError:
                opts[name] = _arh_mes_default()
            continue
        if choices == NUMBER:
            value = (params.get(name) or "").strip()
            opts[name] = value if value.isdigit() else ""
            continue
        value = params.get(name, choices[0][0])
        opts[name] = value if value in {c for c, _ in choices} else choices[0][0]
    return rdef.build(city, opts), opts


def to_xlsx(report, out):
    """Выгрузка отчета в Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Отчет"
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)
    width = max(len(t.cols) for t in report.tables)
    r = 0
    import textwrap
    for line in report.approve:
        for part in textwrap.wrap(line, 45) or [""]:
            r += 1
            ws.cell(r, max(1, width - 2), part)
    mark = PatternFill("solid", fgColor="DCE6F2")
    # Заголовок отчета — по середине над таблицей.
    for k, line in enumerate([report.title] + list(report.lines)):
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)
        cell = ws.cell(r, 1, line)
        cell.alignment = Alignment(horizontal="center")
        if k == 0:
            cell.font = Font(bold=True, size=13)
    widths = {}

    def put(row, values, cols, font=None, marked=False, center=False):
        if marked:
            ws.row_dimensions[row].height = 24
        for i, (v, c) in enumerate(zip(values, cols), 1):
            cell = ws.cell(row, i, v)
            cell.border = box
            if font:
                cell.font = font
            if marked:
                cell.fill, cell.font = mark, bold
                cell.alignment = Alignment(vertical="center", horizontal="center" if center and i > 1 else None)
            if c.kind == "num" and v is not None and v != "":
                cell.number_format = "0." + "0" * c.places if c.places else "0"
            elif c.kind == "date" and v:
                cell.number_format = "DD.MM.YYYY"
            if c.kind == "text":
                widths[i] = max(widths.get(i, 8), min(60, len(str(v or "")) + 6))
            else:
                widths[i] = max(widths.get(i, 8), 14)

    for t in report.tables:
        r += 2
        if t.title:
            ws.cell(r, 1, t.title).font = bold
            r += 1
        if t.head:
            busy = set()
            for k, hrow in enumerate(t.head):
                col = 1
                for title, cs, rs in hrow:
                    while (r + k, col) in busy:
                        col += 1
                    if cs > 1 or rs > 1:
                        ws.merge_cells(start_row=r + k, start_column=col, end_row=r + k + rs - 1,
                                       end_column=col + cs - 1)
                    for dr in range(rs):
                        for dc in range(cs):
                            busy.add((r + k + dr, col + dc))
                            ws.cell(r + k + dr, col + dc).border = box
                    cell = ws.cell(r + k, col, title)
                    cell.font = bold
                    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
                    col += cs
            r += len(t.head) - 1
        elif t.groups and any(g for g, _ in t.groups):
            col = 1
            for g, span in t.groups:
                if span > 1:
                    ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + span - 1)
                cell = ws.cell(r, col, g)
                cell.font, cell.alignment, cell.border = bold, Alignment(horizontal="center"), box
                col += span
            r += 1
        for i, c in enumerate(t.cols, 1):
            if t.head:
                break
            cell = ws.cell(r, i, c.title)
            cell.font, cell.border = bold, box
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        for s in t.sections:
            if s.title:
                r += 1
                ws.cell(r, 1, s.title).font = bold
            for row in s.rows:
                r += 1
                put(r, row, t.cols, marked=s.mark_rows, center=True)
            if s.foot:
                r += 1
                put(r, s.foot, t.cols, bold, marked=s.mark_foot)
        if t.total:
            r += 1
            put(r, t.total, t.cols, bold)
    for note in report.notes:
        r += 2
        ws.cell(r, 1, note).font = bold
    for sig in report.signatures:
        r += 2
        ws.cell(r, 2, sig)
    for i in range(1, width + 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(i, 12)
    if report.landscape:
        ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    wb.save(out)
